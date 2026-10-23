import hashlib
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from m2_identity import new_id
from m3_core import qmd_body_sha256, write_custom_properties
from m3_transaction import FileTransaction
from importlib.machinery import SourceFileLoader

import importlib.util
spec = importlib.util.spec_from_loader("m3_importer", SourceFileLoader("m3_importer", str(ROOT / "scripts" / "import-pptx.py")))
importer = importlib.util.module_from_spec(spec)
spec.loader.exec_module(importer)

from PIL import Image
from pptx import Presentation
from pptx.util import Inches


DOC = "document-123e4567-e89b-42d3-a456-426614174000"
SID = "slide-123e4567-e89b-42d3-a456-426614174001"
EID = "element-123e4567-e89b-42d3-a456-426614174002"


def make_deck(root: Path, *, extra_textbox=False, caption=True, title_marker=False, crop=None):
    qmd = root / "deck.qmd"
    qmd.write_bytes(b"\xef\xbb\xbf---\r\ntitle: Deck\r\ncustom: keep\r\n---\r\n\n## Old {slide-id=\"" + SID.encode() + b"\"}\n\nOriginal text\n")
    props = Presentation()
    slide = props.slides.add_slide(props.slide_layouts[1])
    slide.name = SID
    title_shape = slide.shapes.title
    title_shape.text = "Edited title"
    body = slide.placeholders[1]
    tf = body.text_frame
    tf.clear()
    tf.paragraphs[0].text = "Ordinary paragraph"
    p = tf.add_paragraph()
    p.text = "Bullet item"
    pp = p._p.get_or_add_pPr()
    pp.append(pp.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar", {"char": "•"}))
    p2 = tf.add_paragraph()
    p2.text = "Numbered item"
    pp2 = p2._p.get_or_add_pPr()
    pp2.append(pp2.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}buAutoNum", {"type": "arabicPeriod"}))
    img = root / "source.png"
    Image.new("RGB", (32, 24), "red").save(img)
    picture = slide.shapes.add_picture(str(img), Inches(1), Inches(3), Inches(2), Inches(1.5))
    picture.name = EID
    if crop:
        fill = picture._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}blipFill")
        attrs = {"left": "l", "top": "t", "right": "r", "bottom": "b"}
        src = fill.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}srcRect", {attrs[k]: str(v) for k, v in crop.items()})
        fill.append(src)
    if caption:
        caption_box = slide.shapes.add_textbox(Inches(1), Inches(4.6), Inches(2), Inches(.3))
        caption_box.name = f"mdst-caption-{EID}"
        caption_box.text = "A caption"
    if extra_textbox:
        slide.shapes.add_textbox(Inches(1), Inches(5), Inches(2), Inches(.3)).text = "Unsupported"
    if title_marker:
        title_slide = props.slides.add_slide(props.slide_layouts[0])
        title_slide.name = f"mdst-title-{DOC.removeprefix('document-')}"
        title_slide.shapes.title.text = "Deck"
        # Match the T-301 Quarto render output: only the generated title
        # placeholder is local; the unused layout subtitle stays inherited.
        for shape in list(title_slide.shapes):
            if shape.is_placeholder and int(shape.placeholder_format.type) == 4:
                shape._element.getparent().remove(shape._element)
    pptx = root / "deck.pptx"
    props.save(pptx)
    write_custom_properties(pptx, {
        "md-slide-tool.managed": "1",
        "md-slide-tool.format-version": "1",
        "md-slide-tool.document-id": DOC,
        "md-slide-tool.qmd-body-sha256": qmd_body_sha256(qmd.read_bytes()),
    })
    assets = root / "deck.assets.json"
    assets.write_text(json.dumps({"schemaVersion": 2, "documentId": DOC,
        "slideSize": {"width": int(props.slide_width), "height": int(props.slide_height), "unit": "emu"},
        "elements": [{"slideId": SID, "elementId": EID, "type": "image", "path": "source.png",
                      "x": int(picture.left), "y": int(picture.top), "width": int(picture.width), "height": int(picture.height)}]}), encoding="utf-8")
    return pptx, qmd, assets


def do_import(pptx, qmd, assets):
    props = importer.read_custom_properties(pptx)
    tx = FileTransaction(qmd.parent, props["md-slide-tool.document-id"])
    result = importer.reconstruct(pptx, qmd, assets, tx)
    tx.commit()
    return result


class M3ImportTests(unittest.TestCase):
    def test_m3_render_metadata_title_caption_and_crop_roundtrip(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            qmd = root / "rendered.qmd"
            qmd.write_text(f'---\ntitle: Deck\n---\n\n## Normal {{slide-id="{SID}"}}\n\n![Caption](source.png){{element-id="{EID}"}}\n', encoding="utf-8")
            Image.new("RGB", (50, 40), "blue").save(root / "source.png")
            images_json = root / ".rendered.images.json"
            images_json.write_text(json.dumps([{"title": "Normal", "slide_id": SID, "images": [{"path": "source.png", "element_id": EID, "caption": "Caption"}]}]), encoding="utf-8")
            pptx = root / "rendered.pptx"
            prs = Presentation()
            title = prs.slides.add_slide(prs.slide_layouts[0])
            title.shapes.title.text = "Deck"
            normal = prs.slides.add_slide(prs.slide_layouts[1])
            normal.shapes.title.text = "Normal"
            prs.save(pptx)
            command = [sys.executable, str(ROOT / "scripts" / "postprocess-pptx.py"), "--pptx", str(pptx), "--images-json", str(images_json), "--assets", str(root / "rendered.assets.json"), "--qmd", str(qmd)]
            result = subprocess.run(command, text=True, capture_output=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            from m3_core import read_custom_properties
            props = read_custom_properties(pptx)
            self.assertEqual(props["md-slide-tool.managed"], "1")
            self.assertEqual(props["md-slide-tool.qmd-body-sha256"], qmd_body_sha256(qmd.read_bytes()))
            data = json.loads((root / "rendered.assets.json").read_text(encoding="utf-8"))
            self.assertEqual(data["schemaVersion"], 2)
            self.assertTrue(data["documentId"].startswith("document-"))
            reopened = Presentation(str(pptx))
            self.assertEqual(reopened.slides[0].name, f"mdst-title-{data['documentId'].removeprefix('document-')}")
            caption_shape = next(s for s in reopened.slides[1].shapes if s.shape_type == 17)
            self.assertEqual(caption_shape.name, f"mdst-caption-{EID}")
            data["elements"][0]["crop"] = {"unit": "ooxml-percent", "left": -5000, "top": 0, "right": 10000, "bottom": 0}
            data["elements"][0]["x"] = -5000
            (root / "rendered.assets.json").write_text(json.dumps(data), encoding="utf-8")
            prs = Presentation()
            prs.slides.add_slide(prs.slide_layouts[0]).shapes.title.text = "Deck"
            prs.slides.add_slide(prs.slide_layouts[1]).shapes.title.text = "Normal"
            prs.save(pptx)
            result = subprocess.run(command, text=True, capture_output=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            reopened = Presentation(str(pptx))
            picture = next(s for s in reopened.slides[1].shapes if s.shape_type == 13)
            src_rect = picture._element.xpath("./p:blipFill/a:srcRect")[0]
            self.assertEqual(src_rect.get("l"), "-5000")
            self.assertEqual(int(picture.left), -5000)
            fill_children = [child.tag.rsplit("}", 1)[-1] for child in picture._element.xpath("./p:blipFill")[0]]
            self.assertLess(fill_children.index("srcRect"), fill_children.index("stretch"))

    def test_import_escapes_literal_pptx_text_from_markdown_structure(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            pptx, qmd, assets = make_deck(root, caption=False)
            prs = Presentation(str(pptx))
            slide = next(s for s in prs.slides if s.name == SID)
            slide.shapes.title.text = "# Heading <img>"
            body = next(s for s in slide.shapes if s.is_placeholder and int(s.placeholder_format.type) in (2, 7))
            body.text_frame.paragraphs[0].text = "![literal](evil.png)"
            body.text_frame.paragraphs[1].text = "- nested"
            body.text_frame.paragraphs[2].text = "1. numbered"
            prs.save(pptx)
            write_custom_properties(pptx, {
                "md-slide-tool.qmd-body-sha256": qmd_body_sha256(qmd.read_bytes())
            })

            do_import(pptx, qmd, assets)
            body = qmd.read_bytes().decode("utf-8")
            self.assertIn(r"## \# Heading \<img\> {slide-id=", body)
            self.assertIn(r"\!\[literal\]\(evil\.png\)", body)
            self.assertIn(r"- \- nested", body)
            self.assertIn(r"1. 1\. numbered", body)
            self.assertNotIn("![literal](evil.png)", body)
            self.assertNotIn("<img>", body)

    def test_t301_fixture_confirms_inherited_bullet_classification(self):
        prs = Presentation(str(ROOT / "tests" / "fixtures" / "m3" / "structure.pptx"))
        content = next(shape for slide in prs.slides for shape in slide.shapes
                       if shape.is_placeholder and int(shape.placeholder_format.type) == 7)
        self.assertEqual([kind for kind, _ in importer.text_and_kind(content)],
                         ["paragraph", "bullet", "bullet", "numbered", "numbered"])

    def test_supported_title_layouts_with_lost_markers_are_ambiguous_even_with_slide_ids(self):
        # The checked-in T-301 templates provide CENTER_TITLE + SUBTITLE.
        for fixture in ("structure.pptx", "template.pptx"):
            template = Presentation(str(ROOT / "tests" / "fixtures" / "m3" / fixture))
            title_layout = next(layout for layout in template.slide_layouts
                                if {int(shape.placeholder_format.type) for shape in layout.shapes
                                    if shape.is_placeholder} >= {3, 4})
            self.assertTrue(importer.is_title_layout(type("Slide", (), {"slide_layout": title_layout})()))

        # Also exercise the structurally equivalent TITLE + SUBTITLE variant.
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            pptx, qmd, assets = make_deck(root, caption=False, title_marker=True)
            prs = Presentation(str(pptx))
            title_slide = prs.slides[-1]
            layout = title_slide.slide_layout
            center_title = next(shape for shape in layout.shapes
                                if shape.is_placeholder and int(shape.placeholder_format.type) == 3)
            center_title._element.nvSpPr.nvPr.ph.set("type", "title")
            title_slide.name = new_id("slide")
            prs.save(pptx)
            write_custom_properties(pptx, {"md-slide-tool.qmd-body-sha256": qmd_body_sha256(qmd.read_bytes())})
            with self.assertRaisesRegex(ValueError, "title-layout slide is ambiguous"):
                do_import(pptx, qmd, assets)

    def test_title_paragraph_bullet_numbered_and_caption_import_preserve_front_matter(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            pptx, qmd, assets = make_deck(root, crop={"left": "-5000", "top": "0", "right": "10000", "bottom": "0"})
            before, _ = importer.split_front_matter(qmd.read_bytes())
            do_import(pptx, qmd, assets)
            updated = qmd.read_bytes()
            self.assertTrue(updated.startswith(before))
            self.assertIn(b"## Edited title {slide-id=\"" + SID.encode() + b"\"}", updated)
            self.assertIn(b"Ordinary paragraph", updated)
            self.assertIn(b"- Bullet item", updated)
            self.assertIn(b"1. Numbered item", updated)
            self.assertIn(b"![A caption](source.png)", updated)
            data = json.loads(assets.read_text(encoding="utf-8"))
            self.assertEqual(data["schemaVersion"], 2)
            self.assertEqual(data["elements"][0]["crop"]["left"], -5000)
            self.assertEqual(data["elements"][0]["x"], 914400)

    def test_unsupported_textbox_fails_before_any_target_changes(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            pptx, qmd, assets = make_deck(root, extra_textbox=True)
            originals = (pptx.read_bytes(), qmd.read_bytes(), assets.read_bytes())
            with self.assertRaisesRegex(ValueError, "unsupported slide-local shape"):
                do_import(pptx, qmd, assets)
            self.assertEqual((pptx.read_bytes(), qmd.read_bytes(), assets.read_bytes()), originals)

    def test_baseline_mismatch_fails_before_transaction_staging(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            pptx, qmd, assets = make_deck(root)
            qmd.write_bytes(qmd.read_bytes() + b"later edit")
            with self.assertRaisesRegex(ValueError, "QMD_CHANGED_SINCE_RENDER"):
                do_import(pptx, qmd, assets)
            self.assertFalse((root / ".md-slide-tool" / "transactions").exists())

    def test_unmarked_pptx_and_document_pair_mismatch_fail(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            pptx, qmd, assets = make_deck(root)
            write_custom_properties(pptx, {"md-slide-tool.managed": "0"})
            with self.assertRaisesRegex(ValueError, "unmanaged"):
                do_import(pptx, qmd, assets)

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            pptx, qmd, assets = make_deck(root)
            data = json.loads(assets.read_text(encoding="utf-8"))
            data["documentId"] = "document-123e4567-e89b-42d3-a456-426614174099"
            assets.write_text(json.dumps(data), encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "documentId mismatch"):
                do_import(pptx, qmd, assets)
            self.assertFalse((root / ".md-slide-tool" / "transactions").exists())

    def test_new_slide_and_picture_receive_ids_and_extracted_image(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            pptx, qmd, assets = make_deck(root, caption=False)
            prs = Presentation(str(pptx))
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.name = ""
            slide.shapes.title.text = "New slide"
            newer = root / "new.png"
            Image.new("RGB", (20, 20), "green").save(newer)
            slide.shapes.add_picture(str(newer), Inches(2), Inches(2), Inches(1), Inches(1))
            prs.save(pptx)
            write_custom_properties(pptx, {"md-slide-tool.qmd-body-sha256": qmd_body_sha256(qmd.read_bytes())})
            do_import(pptx, qmd, assets)
            data = json.loads(assets.read_text(encoding="utf-8"))
            self.assertEqual(len(data["elements"]), 2)
            new_entry = next(item for item in data["elements"] if item["slideId"] != SID)
            self.assertTrue(new_entry["slideId"].startswith("slide-"))
            self.assertTrue(new_entry["elementId"].startswith("element-"))
            extracted = root / new_entry["path"]
            self.assertTrue(new_entry["path"].startswith("img/imported/"))
            self.assertEqual(extracted.read_bytes(), newer.read_bytes())
            reopened = Presentation(str(pptx))
            self.assertEqual(reopened.slides[1].name, new_entry["slideId"])
            self.assertEqual(next(s for s in reopened.slides[1].shapes if s.shape_type == 13).name, new_entry["elementId"])

    def test_same_identified_picture_replacement_keeps_id_and_shared_source(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            pptx, qmd, assets = make_deck(root)
            original_source = (root / "source.png").read_bytes()
            replacement = root / "replacement.png"
            Image.new("RGB", (32, 24), "blue").save(replacement)
            prs = Presentation(str(pptx))
            slide = prs.slides[0]
            picture = next(s for s in slide.shapes if s.shape_type == 13)
            picture._element.getparent().remove(picture._element)
            new_picture = slide.shapes.add_picture(str(replacement), Inches(1), Inches(3), Inches(2), Inches(1.5))
            new_picture.name = EID
            prs.save(pptx)
            write_custom_properties(pptx, {"md-slide-tool.qmd-body-sha256": qmd_body_sha256(qmd.read_bytes())})
            do_import(pptx, qmd, assets)
            data = json.loads(assets.read_text(encoding="utf-8"))
            self.assertEqual(data["elements"][0]["elementId"], EID)
            self.assertEqual(data["elements"][0]["path"], f"img/imported/{EID}.png")
            self.assertEqual((root / "source.png").read_bytes(), original_source)
            self.assertEqual((root / data["elements"][0]["path"]).read_bytes(), replacement.read_bytes())

    def test_deleted_image_is_removed_from_assets_and_qmd(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            pptx, qmd, assets = make_deck(root)
            prs = Presentation(str(pptx))
            for shape in list(prs.slides[0].shapes):
                if shape.shape_type == 13 or shape.name.startswith("mdst-caption-"):
                    shape._element.getparent().remove(shape._element)
            prs.save(pptx)
            write_custom_properties(pptx, {"md-slide-tool.qmd-body-sha256": qmd_body_sha256(qmd.read_bytes())})
            do_import(pptx, qmd, assets)
            self.assertEqual(json.loads(assets.read_text(encoding="utf-8"))["elements"], [])
            self.assertNotIn(b"![", qmd.read_bytes())

    def test_rotated_picture_fails_instead_of_losing_rotation(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            pptx, qmd, assets = make_deck(root, caption=False)
            prs = Presentation(str(pptx))
            picture = next(s for s in prs.slides[0].shapes if s.shape_type == 13)
            picture.rotation = 20
            prs.save(pptx)
            write_custom_properties(pptx, {"md-slide-tool.qmd-body-sha256": qmd_body_sha256(qmd.read_bytes())})
            with self.assertRaisesRegex(ValueError, "picture rotation is unsupported"):
                do_import(pptx, qmd, assets)

    def test_invalid_crop_orphan_and_duplicate_caption_markers_fail_before_staging(self):
        for mode in ("invalid_crop", "orphan_caption", "duplicate_caption"):
            with self.subTest(mode=mode), tempfile.TemporaryDirectory() as temp:
                root = Path(temp)
                pptx, qmd, assets = make_deck(root)
                prs = Presentation(str(pptx))
                slide = prs.slides[0]
                if mode == "invalid_crop":
                    picture = next(s for s in slide.shapes if s.shape_type == 13)
                    fill = picture._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}blipFill")
                    src = fill.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}srcRect", {"l": "60000", "r": "50000"})
                    fill.append(src)
                elif mode == "orphan_caption":
                    box = slide.shapes.add_textbox(Inches(1), Inches(5), Inches(2), Inches(.3))
                    box.name = "mdst-caption-element-123e4567-e89b-42d3-a456-426614174099"
                else:
                    box = slide.shapes.add_textbox(Inches(1), Inches(5), Inches(2), Inches(.3))
                    box.name = f"mdst-caption-{EID}"
                prs.save(pptx)
                write_custom_properties(pptx, {"md-slide-tool.qmd-body-sha256": qmd_body_sha256(qmd.read_bytes())})
                expected = "crop has non-positive" if mode == "invalid_crop" else "orphan caption" if mode == "orphan_caption" else "duplicate caption"
                with self.assertRaisesRegex(ValueError, expected):
                    do_import(pptx, qmd, assets)
                self.assertFalse((root / ".md-slide-tool" / "transactions").exists())

    def test_caption_marker_on_non_textbox_fails_closed(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            pptx, qmd, assets = make_deck(root, caption=False)
            prs = Presentation(str(pptx))
            shape = prs.slides[0].shapes.add_shape(1, Inches(1), Inches(5), Inches(2), Inches(.3))
            shape.name = f"mdst-caption-{EID}"
            shape.text = "A deceptive caption"
            prs.save(pptx)
            write_custom_properties(pptx, {"md-slide-tool.qmd-body-sha256": qmd_body_sha256(qmd.read_bytes())})
            with self.assertRaisesRegex(ValueError, "caption marker must be a non-placeholder text box"):
                do_import(pptx, qmd, assets)

    def test_shared_imported_asset_paths_are_not_overwritten_or_deleted(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            pptx, qmd, assets_path = make_deck(root, caption=False)
            prs = Presentation(str(pptx))
            second = prs.slides[0].shapes.add_picture(str(root / "source.png"), Inches(4), Inches(3), Inches(2), Inches(1.5))
            second.name = "element-123e4567-e89b-42d3-a456-426614174099"
            prs.save(pptx)
            shared = f"img/imported/{EID}.png"
            source_bytes = (root / "source.png").read_bytes()
            shared_path = root / shared
            shared_path.parent.mkdir(parents=True)
            shared_path.write_bytes(source_bytes)
            data = json.loads(assets_path.read_text(encoding="utf-8"))
            data["elements"][0]["path"] = shared
            data["elements"].append({**data["elements"][0], "elementId": second.name})
            assets_path.write_text(json.dumps(data), encoding="utf-8")
            # Change one picture, while the second picture still points at
            # the shared original bytes.
            prs = Presentation(str(pptx))
            first = next(shape for shape in prs.slides[0].shapes if shape.name == EID)
            left, top, width, height = first.left, first.top, first.width, first.height
            first._element.getparent().remove(first._element)
            changed = root / "changed.png"
            Image.new("RGB", (32, 24), "blue").save(changed)
            replacement = prs.slides[0].shapes.add_picture(str(changed), left, top, width, height)
            replacement.name = EID
            prs.save(pptx)
            write_custom_properties(pptx, {"md-slide-tool.qmd-body-sha256": qmd_body_sha256(qmd.read_bytes())})
            do_import(pptx, qmd, assets_path)
            updated = json.loads(assets_path.read_text(encoding="utf-8"))
            by_id = {entry["elementId"]: entry for entry in updated["elements"]}
            self.assertEqual(by_id[second.name]["path"], shared)
            self.assertNotEqual(by_id[EID]["path"], shared)
            self.assertEqual(shared_path.read_bytes(), source_bytes)
            self.assertTrue((root / by_id[EID]["path"]).is_file())

    def test_duplicate_slide_and_element_ids_fail(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            pptx, qmd, assets = make_deck(root, caption=False)
            prs = Presentation(str(pptx))
            duplicate = prs.slides.add_slide(prs.slide_layouts[1])
            duplicate.name = SID
            duplicate.shapes.title.text = "Duplicate slide"
            prs.save(pptx)
            write_custom_properties(pptx, {"md-slide-tool.qmd-body-sha256": qmd_body_sha256(qmd.read_bytes())})
            with self.assertRaisesRegex(ValueError, "duplicate slide-id"):
                do_import(pptx, qmd, assets)

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            pptx, qmd, assets = make_deck(root, caption=False)
            prs = Presentation(str(pptx))
            second = prs.slides[0].shapes.add_picture(str(root / "source.png"), Inches(3), Inches(3), Inches(1), Inches(1))
            second.name = EID
            prs.save(pptx)
            write_custom_properties(pptx, {"md-slide-tool.qmd-body-sha256": qmd_body_sha256(qmd.read_bytes())})
            with self.assertRaisesRegex(ValueError, "duplicate element-id"):
                do_import(pptx, qmd, assets)

    def test_moved_and_deleted_title_role_are_excluded_and_duplicate_or_lost_marker_fails(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            pptx, qmd, assets = make_deck(root, caption=False, title_marker=True)
            do_import(pptx, qmd, assets)
            self.assertNotIn(b"## Deck", qmd.read_bytes())
            self.assertIn(b"## Edited title", qmd.read_bytes())

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            pptx, qmd, assets = make_deck(root, caption=False, title_marker=True)
            prs = Presentation(str(pptx))
            rel_id = prs.slides._sldIdLst[-1].rId
            prs.part.drop_rel(rel_id)
            prs.slides._sldIdLst.remove(prs.slides._sldIdLst[-1])
            prs.save(pptx)
            write_custom_properties(pptx, {"md-slide-tool.qmd-body-sha256": qmd_body_sha256(qmd.read_bytes())})
            do_import(pptx, qmd, assets)
            self.assertIn(b"## Edited title", qmd.read_bytes())

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            pptx, qmd, assets = make_deck(root, caption=False, title_marker=True)
            prs = Presentation(str(pptx))
            duplicate = prs.slides.add_slide(prs.slide_layouts[0])
            duplicate.name = f"mdst-title-{DOC.removeprefix('document-')}"
            prs.save(pptx)
            write_custom_properties(pptx, {"md-slide-tool.qmd-body-sha256": qmd_body_sha256(qmd.read_bytes())})
            with self.assertRaisesRegex(ValueError, "duplicate title-role"):
                do_import(pptx, qmd, assets)

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            pptx, qmd, assets = make_deck(root, caption=False, title_marker=True)
            prs = Presentation(str(pptx))
            prs.slides[-1].name = ""
            prs.save(pptx)
            write_custom_properties(pptx, {"md-slide-tool.qmd-body-sha256": qmd_body_sha256(qmd.read_bytes())})
            with self.assertRaisesRegex(ValueError, "title-layout slide is ambiguous"):
                do_import(pptx, qmd, assets)

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            pptx, qmd, assets = make_deck(root, caption=False, title_marker=True)
            prs = Presentation(str(pptx))
            prs.slides[-1].name = new_id("slide")
            prs.save(pptx)
            write_custom_properties(pptx, {"md-slide-tool.qmd-body-sha256": qmd_body_sha256(qmd.read_bytes())})
            with self.assertRaisesRegex(ValueError, "title-layout slide is ambiguous"):
                do_import(pptx, qmd, assets)

    def test_unsupported_content_on_ignored_title_slide_fails_closed(self):
        for add_picture in (False, True):
            with self.subTest(add_picture=add_picture), tempfile.TemporaryDirectory() as temp:
                root = Path(temp)
                pptx, qmd, assets = make_deck(root, caption=False, title_marker=True)
                prs = Presentation(str(pptx))
                title_slide = prs.slides[-1]
                if add_picture:
                    title_slide.shapes.add_picture(str(root / "source.png"), Inches(1), Inches(1), Inches(1), Inches(1))
                else:
                    title_slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(.3)).text = "User content"
                prs.save(pptx)
                write_custom_properties(pptx, {"md-slide-tool.qmd-body-sha256": qmd_body_sha256(qmd.read_bytes())})
                with self.assertRaisesRegex(ValueError, "unsupported slide-local structure on title slide"):
                    do_import(pptx, qmd, assets)
                self.assertFalse((root / ".md-slide-tool" / "transactions").exists())

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            pptx, qmd, assets = make_deck(root, caption=False, title_marker=True)
            prs = Presentation(str(pptx))
            layout_subtitle = next(shape for shape in prs.slides[-1].slide_layout.shapes
                                   if shape.is_placeholder and int(shape.placeholder_format.type) == 4)
            prs.slides[-1].shapes.clone_placeholder(layout_subtitle)
            prs.save(pptx)
            write_custom_properties(pptx, {"md-slide-tool.qmd-body-sha256": qmd_body_sha256(qmd.read_bytes())})
            with self.assertRaisesRegex(ValueError, "unsupported slide-local structure on title slide"):
                do_import(pptx, qmd, assets)

    def test_all_normal_slides_can_be_deleted_while_presentation_remains_managed(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            pptx, qmd, assets = make_deck(root, caption=False, title_marker=True)
            original_front, _ = importer.split_front_matter(qmd.read_bytes())
            prs = Presentation(str(pptx))
            rel_id = prs.slides._sldIdLst[0].rId
            prs.part.drop_rel(rel_id)
            prs.slides._sldIdLst.remove(prs.slides._sldIdLst[0])
            prs.save(pptx)
            write_custom_properties(pptx, {"md-slide-tool.qmd-body-sha256": qmd_body_sha256(qmd.read_bytes())})
            do_import(pptx, qmd, assets)
            self.assertEqual(qmd.read_bytes(), original_front)
            self.assertEqual(json.loads(assets.read_text(encoding="utf-8"))["elements"], [])
            self.assertEqual(importer.read_custom_properties(pptx)["md-slide-tool.managed"], "1")


if __name__ == "__main__":
    unittest.main()
