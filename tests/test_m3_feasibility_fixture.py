#!/usr/bin/env python3
"""Verify structural facts in the T-301 PPTX rendered through the repo pipeline.

This is a feasibility fixture only. It does not implement M3 import behavior.
"""

import hashlib
import json
import re
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls


ROOT = Path(__file__).resolve().parents[1]
FIXTURE = ROOT / "tests" / "fixtures" / "m3"
TEMPLATE_ROOT = ROOT / "templates" / "pptx"
TEMPLATE = ROOT / "templates" / "pptx" / "template" / "template.pptx"
SELFINTRO_DIR = FIXTURE / "selfintroduction"
SELFINTRO_SOURCE_TEMPLATE = ROOT / "projects" / "selfintroduction" / "template.pptx"
SELFINTRO_TEMPLATE_COPY = SELFINTRO_DIR / "template.source.pptx"
SELFINTRO_REFERENCE_DOC = SELFINTRO_DIR / "template.pptx"
SELFINTRO_RENDERED = SELFINTRO_DIR / "structure.pptx"
TEMPLATE_COPY = FIXTURE / "template.source.pptx"
RENDERED = FIXTURE / "structure.pptx"
NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "ct": "http://schemas.openxmlformats.org/package/2006/content-types",
    "cp": "http://schemas.openxmlformats.org/officeDocument/2006/custom-properties",
    "vt": "http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes",
}


def require(condition, message):
    if not condition:
        raise AssertionError(message)


def custom_properties_xml():
    return f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Properties xmlns="{NS['cp']}" xmlns:vt="{NS['vt']}">
  <property fmtid="{{D5CDD505-2E9C-101B-9397-08002B2CF9AE}}" pid="2" name="md-slide-tool.managed"><vt:lpwstr>1</vt:lpwstr></property>
</Properties>'''


def add_custom_properties(package_path):
    """Add a standard custom-properties part without a python-pptx API."""
    temp_path = package_path.with_suffix(".rewrite.pptx")
    with zipfile.ZipFile(package_path, "r") as source, zipfile.ZipFile(temp_path, "w", zipfile.ZIP_DEFLATED) as target:
        names = set(source.namelist())
        content_types = ET.fromstring(source.read("[Content_Types].xml"))
        if not any(node.attrib.get("PartName") == "/docProps/custom.xml" for node in content_types):
            ET.SubElement(content_types, f"{{{NS['ct']}}}Override", {
                "PartName": "/docProps/custom.xml",
                "ContentType": "application/vnd.openxmlformats-officedocument.custom-properties+xml",
            })
        relationships = ET.fromstring(source.read("_rels/.rels"))
        rel_type = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties"
        if not any(node.attrib.get("Type") == rel_type for node in relationships):
            used = {node.attrib.get("Id") for node in relationships}
            index = 1
            while f"rId{index}" in used:
                index += 1
            ET.SubElement(relationships, f"{{{NS['rel']}}}Relationship", {
                "Id": f"rId{index}", "Type": rel_type, "Target": "docProps/custom.xml",
            })
        for item in source.infolist():
            if item.filename in {"[Content_Types].xml", "_rels/.rels", "docProps/custom.xml"}:
                continue
            target.writestr(item, source.read(item.filename))
        target.writestr("[Content_Types].xml", ET.tostring(content_types, encoding="utf-8", xml_declaration=True))
        target.writestr("_rels/.rels", ET.tostring(relationships, encoding="utf-8", xml_declaration=True))
        target.writestr("docProps/custom.xml", custom_properties_xml())
    temp_path.replace(package_path)


def raw_crop(shape):
    node = shape._element.blipFill.find("{http://schemas.openxmlformats.org/drawingml/2006/main}srcRect")
    return None if node is None else {key: int(value) for key, value in node.attrib.items()}


def set_raw_crop(shape, values):
    blip_fill = shape._element.blipFill
    src_rect = blip_fill.find("{http://schemas.openxmlformats.org/drawingml/2006/main}srcRect")
    if src_rect is None:
        src_rect = parse_xml(f"<a:srcRect {nsdecls('a')}/>")
        blip_fill.insert(1, src_rect)
    for key, value in values.items():
        src_rect.set(key, str(value))


def shape_inventory(shape):
    placeholder = None
    if shape.is_placeholder:
        placeholder = {
            "type": int(shape.placeholder_format.type),
            "index": int(shape.placeholder_format.idx),
        }
    return {
        "id": int(shape.shape_id),
        "name": shape.name,
        "shapeType": str(shape.shape_type),
        "placeholder": placeholder,
        "geometryEmu": {
            "left": int(shape.left), "top": int(shape.top),
            "width": int(shape.width), "height": int(shape.height),
        },
        "text": shape.text if shape.has_text_frame else None,
        "xml": shape._element.xml,
    }


def classify_local_shape(slide, shape):
    """Structural M3 feasibility rules; unmarked text boxes stay unsupported."""
    if shape.is_placeholder:
        placeholder_type = int(shape.placeholder_format.type)
        placeholder_index = int(shape.placeholder_format.idx)
        if placeholder_type == 3 and placeholder_index == 0:
            return "center-title-placeholder"
        if placeholder_type == 1 and placeholder_index == 0:
            return "normal-slide-title-placeholder"
        if placeholder_type == 7 and placeholder_index == 1:
            return "normal-slide-body-placeholder"
        return "unsupported-placeholder"
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return "picture"
    if shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX and shape.name.startswith("mdst-caption-"):
        return "caption-companion"
    if shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
        return "unsupported-unmarked-textbox"
    return "unsupported-slide-local-shape"


def template_inventory(path):
    prs = Presentation(str(path))
    return {
        "path": path.relative_to(ROOT).as_posix(),
        "slideCount": len(prs.slides),
        "masterCount": len(prs.slide_masters),
        "layoutCount": len(prs.slide_layouts),
        "slides": [{
            "name": slide.name, "layout": slide.slide_layout.name,
            "shapes": [shape_inventory(shape) for shape in slide.shapes],
        } for slide in prs.slides],
        "masters": [{
            "name": master.name,
            "shapes": [{
                **shape_inventory(shape),
                "structuralClass": "inherited-master-placeholder" if shape.is_placeholder else "inherited-master-object",
            } for shape in master.shapes],
        } for master in prs.slide_masters],
        "layouts": [{
            "name": layout.name, "master": layout.slide_master.name,
            "shapes": [{
                **shape_inventory(shape),
                "structuralClass": "inherited-layout-placeholder" if shape.is_placeholder else "inherited-layout-object",
            } for shape in layout.shapes],
        } for layout in prs.slide_layouts],
    }


def verify_selfintroduction_fixture():
    require(SELFINTRO_RENDERED.is_file(), f"Missing self-introduction rendered fixture: {SELFINTRO_RENDERED}")
    provenance = json.loads((SELFINTRO_DIR / "template-provenance.json").read_text(encoding="utf-8"))
    render_provenance = json.loads((SELFINTRO_DIR / "render-provenance.json").read_text(encoding="utf-8"))
    source_sha256 = hashlib.sha256(SELFINTRO_SOURCE_TEMPLATE.read_bytes()).hexdigest()
    preserved_sha256 = hashlib.sha256(SELFINTRO_TEMPLATE_COPY.read_bytes()).hexdigest()
    reference_sha256 = hashlib.sha256(SELFINTRO_REFERENCE_DOC.read_bytes()).hexdigest()
    output_sha256 = hashlib.sha256(SELFINTRO_RENDERED.read_bytes()).hexdigest()
    qmd_sha256 = hashlib.sha256((SELFINTRO_DIR / "structure.qmd").read_bytes()).hexdigest()
    require(render_provenance["renderExitCode"] == 0, "self-introduction Dev Container render did not record success")
    require(render_provenance["devContainerImageId"] == "sha256:70da969238ff921f30dbc050dbbff38e5dafda2bfdaa7759446966d640f2767b",
            "self-introduction fixture Dev Container image identity changed")
    require(render_provenance["quartoVersion"] == "1.9.36" and render_provenance["pythonPptxVersion"] == "1.0.2",
            "self-introduction fixture render toolchain changed")
    require(render_provenance["output"] == SELFINTRO_RENDERED.relative_to(ROOT).as_posix(), "self-introduction render output path changed")
    require(render_provenance["outputSha256"] == output_sha256 and render_provenance["inputQmdSha256"] == qmd_sha256,
            "self-introduction fixture output/QMD no longer matches render provenance")
    require(provenance["byteIdenticalBeforeNormalization"] is True, "self-introduction template provenance did not record an exact pre-normalization copy")
    require(provenance["sourceSha256"] == source_sha256 == preserved_sha256, "preserved self-introduction source template differs from repository source")
    require(provenance["referenceDocSha256BeforeNormalization"] == source_sha256, "sibling reference-doc was not recorded as byte-identical before normalization")
    require(provenance["sourceTemplate"] == SELFINTRO_SOURCE_TEMPLATE.relative_to(ROOT).as_posix(), "self-introduction template provenance source path changed")
    require(provenance["preservedSourceCopy"] == SELFINTRO_TEMPLATE_COPY.relative_to(ROOT).as_posix(), "preserved template provenance path changed")
    require(provenance["referenceDocBeforeNormalization"] == SELFINTRO_REFERENCE_DOC.relative_to(ROOT).as_posix(), "reference-doc provenance path changed")

    source_template = Presentation(str(SELFINTRO_TEMPLATE_COPY))
    rendered = Presentation(str(SELFINTRO_RENDERED))
    require(len(source_template.slide_masters) == 1 and len(source_template.slide_layouts) == 17, "self-introduction source template inventory changed")
    require(len(rendered.slides) == 2, "self-introduction render must contain generated title and normal slide")
    title_slide, normal_slide = rendered.slides[0], rendered.slides[1]
    require(title_slide.slide_layout.name == "Title Slide", "self-introduction title slide layout changed")
    require(normal_slide.slide_layout.name == "Title and Content", "self-introduction normal slide layout changed")
    title_placeholder = next(shape for shape in title_slide.shapes if shape.is_placeholder)
    normal_title = next(shape for shape in normal_slide.shapes if shape.is_placeholder and int(shape.placeholder_format.type) == 1)
    body = next(shape for shape in normal_slide.shapes if shape.is_placeholder and int(shape.placeholder_format.type) == 7)
    picture = next(shape for shape in normal_slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
    caption = next(shape for shape in normal_slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX)
    require((int(title_placeholder.placeholder_format.type), int(title_placeholder.placeholder_format.idx)) == (3, 0), "self-introduction title placeholder changed")
    require((int(normal_title.placeholder_format.type), int(normal_title.placeholder_format.idx)) == (1, 0), "self-introduction normal title placeholder changed")
    require((int(body.placeholder_format.type), int(body.placeholder_format.idx)) == (7, 1), "self-introduction body placeholder changed")
    require(title_placeholder.text == "M3 self-introduction template fixture", "self-introduction front-matter title content changed")
    require(normal_title.text == "Structural fixture slide", "self-introduction normal title content changed")
    expected_text = [
        "Ordinary paragraph in a supported body placeholder.",
        "Bullet item one", "Bullet item two", "Numbered item one", "Numbered item two",
    ]
    require([para.text for para in body.text_frame.paragraphs] == expected_text, "self-introduction body text/list structure changed")
    pprs = [ET.fromstring(para._p.xml).find("a:pPr", NS) for para in body.text_frame.paragraphs]
    require(pprs[0] is not None and pprs[0].find("a:buNone", NS) is not None, "self-introduction ordinary paragraph must disable bullets")
    require(all(pprs[index] is not None and pprs[index].find("a:buAutoNum", NS) is not None for index in (3, 4)), "self-introduction ordered-list paragraphs must use buAutoNum")
    master_body_style = rendered.slide_masters[0]._element.find(
        ".//{http://schemas.openxmlformats.org/presentationml/2006/main}txStyles/"
        "{http://schemas.openxmlformats.org/presentationml/2006/main}bodyStyle"
    )
    level_one = master_body_style.find("{http://schemas.openxmlformats.org/drawingml/2006/main}lvl1pPr") if master_body_style is not None else None
    require(all(pprs[index] is not None and pprs[index].find("a:buNone", NS) is None and pprs[index].find("a:buChar", NS) is None for index in (1, 2)), "self-introduction bullets must inherit their glyph")
    require(level_one is not None and level_one.find("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar") is not None, "self-introduction bullets must inherit a glyph from master bodyStyle")
    require(picture.image.ext == "png" and picture.image.content_type == "image/png" and len(picture.image.blob) == 68,
            "self-introduction image bytes/format changed")
    require(caption.text == "Caption fixture" and not caption.is_placeholder, "self-introduction renderer caption structure changed")

    qmd_text = (SELFINTRO_DIR / "structure.qmd").read_text(encoding="utf-8")
    assets = json.loads((SELFINTRO_DIR / "structure.assets.json").read_text(encoding="utf-8"))
    slide_match = re.search(r'slide-id="(slide-[0-9a-f-]{36})"', qmd_text)
    element_match = re.search(r'element-id="(element-[0-9a-f-]{36})"', qmd_text)
    require(slide_match is not None and normal_slide.name == slide_match.group(1), "self-introduction rendered slide stable ID does not match QMD")
    require(element_match is not None and picture.name == element_match.group(1), "self-introduction picture element ID does not match QMD")
    require(assets["schemaVersion"] == 1 and len(assets["elements"]) == 1, "self-introduction assets fixture changed")
    require(assets["elements"][0]["slideId"] == normal_slide.name and assets["elements"][0]["elementId"] == picture.name,
            "self-introduction stable IDs do not match assets metadata")
    categories = [[classify_local_shape(slide, shape) for shape in slide.shapes] for slide in rendered.slides]
    require(categories == [
        ["center-title-placeholder"],
        ["normal-slide-title-placeholder", "normal-slide-body-placeholder", "picture", "unsupported-unmarked-textbox"],
    ], f"self-introduction slide-local structural classification changed: {categories}")
    source_inventory = template_inventory(SELFINTRO_TEMPLATE_COPY)
    inherited_layout_objects = [
        {"layout": layout["name"], "shapeType": shape["shapeType"], "structuralClass": shape["structuralClass"]}
        for layout in source_inventory["layouts"] for shape in layout["shapes"]
        if shape["structuralClass"] == "inherited-layout-object"
    ]
    require(inherited_layout_objects, "self-introduction template has no structurally inventoried inherited layout objects")
    require(all(shape["structuralClass"] == "inherited-layout-object" for shape in inherited_layout_objects),
            "layout decorations must be classified by their layout-tree origin")

    return {
        "fixture": SELFINTRO_RENDERED.relative_to(ROOT).as_posix(),
        "templateProvenance": {
            **provenance,
            "sourceSha256Verified": source_sha256,
            "preservedCopySha256Verified": preserved_sha256,
            "normalizedReferenceDocSha256AfterRender": reference_sha256,
            "referenceDocUnchangedByRender": reference_sha256 == source_sha256,
        },
        "renderProvenance": {
            **render_provenance,
            "verifiedOutputSha256": output_sha256,
            "verifiedInputQmdSha256": qmd_sha256,
        },
        "sourceTemplate": source_inventory,
        "inheritedTemplateContent": {
            "layoutObjectCount": len(inherited_layout_objects),
            "layoutObjects": inherited_layout_objects,
            "rule": "Objects enumerated in a layout/master XML tree are inherited template content; only slide.shapes enters slide-local import classification.",
        },
        "rendered": {
            "slideCount": len(rendered.slides),
            "slides": [{
                "name": slide.name, "layout": slide.slide_layout.name,
                "classification": [classify_local_shape(slide, shape) for shape in slide.shapes],
                "shapes": [shape_inventory(shape) for shape in slide.shapes],
            } for slide in rendered.slides],
            "bodyParagraphs": [{"text": para.text, "xml": para._p.xml} for para in body.text_frame.paragraphs],
            "picture": {"name": picture.name, "format": picture.image.ext, "contentType": picture.image.content_type,
                        "byteLength": len(picture.image.blob), "sha256": hashlib.sha256(picture.image.blob).hexdigest()},
            "caption": {"shapeType": str(caption.shape_type), "placeholder": caption.is_placeholder,
                        "name": caption.name, "text": caption.text},
        },
        "classificationRules": {
            "titleRole": "reserved slide-name marker mdst-title-<document-uuid>; placeholder type does not determine slide role",
            "placeholderContent": "PPTX placeholder type/index",
            "picture": "local shape type PICTURE",
            "caption": "local TEXT_BOX only when name starts mdst-caption-; unmarked text boxes are unsupported",
            "masterAndLayoutDecorations": "inherited by relationship to master/layout; excluded from slide.shapes local-content classification",
            "otherSlideLocalShapes": "unsupported and must fail import",
        },
        "limitations": [
            "Quarto emitted warnings that Content with Caption and Blank layouts were absent and fell back to bundled layouts.",
            "Current baseline renderer caption remains unmarked and therefore classifies as unsupported until an M3-capable render stamps its reserved marker.",
        ],
    }


def main():
    require(RENDERED.is_file(), f"Missing rendered fixture: {RENDERED}")
    root_provenance = json.loads((FIXTURE / "template-provenance.json").read_text(encoding="utf-8"))
    root_source_sha256 = hashlib.sha256(TEMPLATE.read_bytes()).hexdigest()
    root_copy_sha256 = hashlib.sha256(TEMPLATE_COPY.read_bytes()).hexdigest()
    root_reference_sha256 = hashlib.sha256((FIXTURE / "template.pptx").read_bytes()).hexdigest()
    require(root_provenance["template-source-sha256"] == root_source_sha256 == root_copy_sha256,
            "root-template source copy does not match the current repository template")
    require(root_provenance["preserved-copy-sha256"] == root_copy_sha256,
            "root-template preserved-source hash does not match the copy")
    require(root_provenance["rendered-sibling-sha256"] == root_reference_sha256,
            "root-template rendered sibling hash does not match the recorded post-render reference-doc")
    require(root_provenance["preservedSourceMatchesRepositorySource"] is True,
            "root-template provenance did not record source-copy equality")
    selfintroduction_fixture = verify_selfintroduction_fixture()

    template = Presentation(str(TEMPLATE))
    rendered = Presentation(str(RENDERED))
    require(len(template.slide_masters) == 1 and len(template.slide_layouts) == 11, "supported template inventory changed")
    template_local = [
        (shape.shape_type, shape.is_placeholder, shape.placeholder_format.type if shape.is_placeholder else None)
        for shape in template.slides[0].shapes
    ]
    require(template_local == [(MSO_SHAPE_TYPE.PLACEHOLDER, True, 3), (MSO_SHAPE_TYPE.PLACEHOLDER, True, 4)], "template slide local placeholder structure changed")

    require(len(rendered.slides) == 3, "fixture expected title, two normal content slides")
    title_slide, content_slide, second_slide = rendered.slides
    require([(shape.placeholder_format.type, shape.placeholder_format.idx) for shape in title_slide.shapes if shape.is_placeholder] == [(3, 0)], "rendered title placeholder changed")
    require([(shape.placeholder_format.type, shape.placeholder_format.idx) for shape in content_slide.shapes if shape.is_placeholder] == [(1, 0), (7, 1)], "rendered content placeholders changed")
    require([(shape.placeholder_format.type, shape.placeholder_format.idx) for shape in second_slide.shapes if shape.is_placeholder] == [(1, 0), (7, 1)], "second content placeholders changed")
    title_shape = next(shape for shape in title_slide.shapes if shape.is_placeholder)
    normal_title_shape = next(shape for shape in content_slide.shapes if shape.is_placeholder and shape.placeholder_format.type == 1)
    require(title_shape.text == "M3 feasibility structure fixture" and len(title_shape.text_frame.paragraphs) == 1,
            "front-matter title slide text/paragraph structure changed")
    require(normal_title_shape.text == "Placeholder and list structures" and len(normal_title_shape.text_frame.paragraphs) == 1,
            "normal slide title text/paragraph structure changed")

    body = next(shape for shape in content_slide.shapes if shape.is_placeholder and shape.placeholder_format.type == 7)
    paragraphs = body.text_frame.paragraphs
    require([p.text for p in paragraphs] == [
        "An ordinary paragraph for the T-301 import feasibility fixture.",
        "Bullet one", "Bullet two", "Numbered one", "Numbered two",
    ], "paragraph text fixture changed")
    p_xml = [ET.fromstring(p._p.xml) for p in paragraphs]
    ppr = [node.find("a:pPr", NS) for node in p_xml]
    require(ppr[0] is not None and ppr[0].find("a:buNone", NS) is not None, "ordinary paragraph must explicitly disable bullets")
    require(all(ppr[i] is not None and ppr[i].find("a:buNone", NS) is None and ppr[i].find("a:buChar", NS) is None for i in (1, 2)), "bullet paragraphs must inherit the content placeholder bullet style")
    require(all(ppr[i] is not None and ppr[i].find("a:buAutoNum", NS) is not None for i in (3, 4)), "ordered list paragraphs must carry a:buAutoNum")
    master_body_style = rendered.slide_masters[0]._element.find(
        ".//{http://schemas.openxmlformats.org/presentationml/2006/main}txStyles/"
        "{http://schemas.openxmlformats.org/presentationml/2006/main}bodyStyle"
    )
    require(master_body_style is not None, "rendered master bodyStyle is missing")
    master_level_one = master_body_style.find("{http://schemas.openxmlformats.org/drawingml/2006/main}lvl1pPr")
    require(master_level_one is not None and master_level_one.find("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar") is not None,
            "bullet paragraphs must inherit a bullet glyph from the rendered master body style")

    picture = next(shape for shape in content_slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
    caption = next(shape for shape in content_slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX)
    require(picture.image.ext == "png" and picture.image.content_type == "image/png", "picture format metadata changed")
    picture_sha256 = hashlib.sha256(picture.image.blob).hexdigest()
    require(len(picture.image.blob) == 68 and picture_sha256 == "defac9092ab8d25febfaf5e4888e518b9c5cffc16e7c89390040a63bd8560ef0", "picture bytes changed")
    require(caption.text == "Caption fixture" and not caption.is_placeholder, "renderer caption textbox structure changed")
    require(not caption.name.startswith("mdst-caption-"), "baseline renderer unexpectedly acquired an M3 caption marker")

    qmd_text = (FIXTURE / "structure.qmd").read_text(encoding="utf-8")
    assets = json.loads((FIXTURE / "structure.assets.json").read_text(encoding="utf-8"))
    slide_ids = [rendered.slides[index].name for index in (1, 2)]
    element_id = picture.name
    require(all(slide_id in qmd_text for slide_id in slide_ids), "rendered slide names do not match stable IDs in the paired QMD")
    require(len(set(slide_ids)) == 2 and all(slide_id.startswith("slide-") for slide_id in slide_ids), "normal slide IDs are missing or duplicated")
    require(element_id.startswith("element-") and element_id in qmd_text, "picture element ID is missing or does not match the paired QMD")
    require(assets["schemaVersion"] == 1 and len(assets["elements"]) == 1, "M2 assets fixture changed")
    asset = assets["elements"][0]
    require(asset["slideId"] == slide_ids[0] and asset["elementId"] == element_id, "PPTX stable IDs do not match assets metadata")

    classification = [
        [classify_local_shape(slide, shape) for shape in slide.shapes]
        for slide in rendered.slides
    ]
    require(classification == [
        ["center-title-placeholder"],
        ["normal-slide-title-placeholder", "normal-slide-body-placeholder", "picture", "unsupported-unmarked-textbox"],
        ["normal-slide-title-placeholder", "normal-slide-body-placeholder"],
    ], f"unexpected slide-local structure/classification: {classification}")

    # A slide collection includes only slide-local objects. Master/layout prototypes are separate.
    inherited_names = {shape.name for shape in content_slide.slide_layout.shapes}
    require(not any(shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX for shape in content_slide.shapes if shape.name in inherited_names), "unexpected inherited object collision")

    with tempfile.TemporaryDirectory(prefix="m3-feasibility-") as temp_dir:
        probe_path = Path(temp_dir) / "probe.pptx"
        probe_path.write_bytes(RENDERED.read_bytes())
        add_custom_properties(probe_path)
        probe = Presentation(str(probe_path))
        title_marker = "mdst-title-00000000-0000-4000-8000-000000000001"
        probe.slides[0].name = title_marker
        probe_content = probe.slides[1]
        probe_picture = next(shape for shape in probe_content.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
        probe_picture.name = element_id
        probe_caption = next(shape for shape in probe_content.shapes if shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX)
        caption_marker = f"mdst-caption-{element_id.removeprefix('element-')}"
        probe_caption.name = caption_marker
        crop = {"l": -5000, "t": 12500, "r": 25000, "b": 0}
        set_raw_crop(probe_picture, crop)
        probe.save(str(probe_path))

        reopened = Presentation(str(probe_path))
        require(reopened.slides[0].name == title_marker, "title-role marker lost on python-pptx reopen")
        require([reopened.slides[i].name for i in (1, 2)] == slide_ids, "normal stable slide IDs lost alongside title marker")
        reopened_picture = next(shape for shape in reopened.slides[1].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
        reopened_caption = next(shape for shape in reopened.slides[1].shapes if shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX)
        require(reopened_picture.name == element_id, "element ID lost alongside markers")
        require(reopened_caption.name == caption_marker, "caption marker lost on python-pptx reopen")
        require(reopened_picture.image.blob == picture.image.blob and reopened_picture.image.ext == "png", "picture bytes/format changed on save/reopen")
        extracted_crop = raw_crop(reopened_picture)
        require(extracted_crop == crop, f"raw crop changed: {extracted_crop}")
        reopened.save(str(probe_path))
        reopened_twice = Presentation(str(probe_path))
        picture_twice = next(shape for shape in reopened_twice.slides[1].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
        require(raw_crop(picture_twice) == extracted_crop, "extracted raw crop could not be reapplied across a second save/reopen")
        require([reopened_twice.slides[index].name for index in (1, 2)] == slide_ids, "stable slide IDs changed after second save/reopen")
        require(next(shape for shape in reopened_twice.slides[1].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE).name == element_id, "element ID changed after second save/reopen")
        require(next(shape for shape in reopened_twice.slides[1].shapes if shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX).name == caption_marker, "caption marker changed after second save/reopen")

        with zipfile.ZipFile(probe_path) as package:
            custom_xml = ET.fromstring(package.read("docProps/custom.xml"))
            prop = custom_xml.find("cp:property[@name='md-slide-tool.managed']/vt:lpwstr", NS)
            require(prop is not None and prop.text == "1", "custom property did not survive python-pptx save/reopen/save")

    report = {
        "status": "PASS",
        "fixture": str(RENDERED.relative_to(ROOT)).replace("\\", "/"),
        "pythonPptxVersion": __import__("pptx").__version__,
        "template": {
            "slideCount": len(template.slides), "masterCount": len(template.slide_masters),
            "layoutCount": len(template.slide_layouts), "localSlideZeroShapes": template_local,
            "slides": [{
                "name": slide.name, "layout": slide.slide_layout.name,
                "shapes": [shape_inventory(shape) for shape in slide.shapes],
            } for slide in template.slides],
            "masters": [{
                "name": master.name,
                "shapes": [shape_inventory(shape) for shape in master.shapes],
            } for master in template.slide_masters],
            "layouts": [{
                "name": layout.name, "master": layout.slide_master.name,
                "shapes": [shape_inventory(shape) for shape in layout.shapes],
            } for layout in template.slide_layouts],
        },
        "repositoryTemplateSources": [
            template_inventory(path)
            for path in sorted(set(TEMPLATE_ROOT.rglob("*.pptx")) | set(ROOT.glob("projects/*/template.pptx")))
        ],
        "defaultTemplateProvenance": {
            **root_provenance,
            "verifiedTemplateSourceSha256": root_source_sha256,
            "verifiedPreservedCopySha256": root_copy_sha256,
            "verifiedRenderedSiblingSha256": root_reference_sha256,
            "renderedSiblingDiffersFromSource": root_reference_sha256 != root_source_sha256,
        },
        "selfIntroductionTemplateRender": selfintroduction_fixture,
        "rendered": {
            "slideCount": len(rendered.slides),
            "slides": [{
                "name": slide.name,
                "layout": slide.slide_layout.name,
                "classification": [classify_local_shape(slide, shape) for shape in slide.shapes],
                "shapes": [shape_inventory(shape) for shape in slide.shapes],
            } for slide in rendered.slides],
            "normalPlaceholderTypesAndIndexes": [[(shape.placeholder_format.type, shape.placeholder_format.idx) for shape in slide.shapes if shape.is_placeholder] for slide in (content_slide, second_slide)],
            "titleParagraphs": {
                "frontMatterTitle": [{"text": para.text, "xml": para._p.xml} for para in title_shape.text_frame.paragraphs],
                "normalSlideTitle": [{"text": para.text, "xml": para._p.xml} for para in normal_title_shape.text_frame.paragraphs],
            },
            "bodyParagraphs": [{"text": para.text, "xml": para._p.xml} for para in paragraphs],
            "caption": {"shapeType": str(caption.shape_type), "placeholder": caption.is_placeholder, "name": caption.name, "text": caption.text},
            "picture": {"name": picture.name, "format": picture.image.ext, "contentType": picture.image.content_type, "byteLength": len(picture.image.blob), "sha256": picture_sha256},
        },
        "saveReopen": {
            "customProperties": "PASS", "titleRoleAndStableIds": "PASS", "captionMarker": "PASS",
            "rawCrop": {"unit": "ooxml-percent", **crop}, "cropTwoSaveReopenCycles": "PASS",
            "titleMarker": title_marker, "normalSlideIds": slide_ids,
            "elementId": element_id, "captionMarker": caption_marker,
        },
        "classificationRules": {
            "titleSlideRole": "reserved slide-name marker mdst-title-<document-uuid>; placeholder type is not used to determine the slide's role",
            "centerTitlePlaceholder": "local placeholder type CENTER_TITLE (3), idx 0",
            "normalTitle": "local placeholder type TITLE (1), idx 0",
            "normalBody": "local placeholder type OBJECT (7), idx 1",
            "picture": "local shape type PICTURE",
            "caption": "local TEXT_BOX only when name starts mdst-caption-; unmarked text boxes are unsupported",
            "otherLocalShapes": "unsupported and must fail import",
        },
        "inheritedVsLocal": {
            "status": "PASS",
            "rule": "Enumerate slide-local shapes only via slide.shapes. Master/layout prototypes are separately inventoried and are not included in the slide-local classification input.",
            "slideClassification": classification,
        },
        "limitations": [
            "T-301 probe verifies python-pptx/Open XML package behavior, not Microsoft PowerPoint desktop save behavior.",
            "Current baseline caption is unmarked and classifies as unsupported-unmarked-textbox; the M3-capable render must stamp the reserved marker before import is implemented.",
            "Bullet style is inherited from the body placeholder/layout; paragraph 0 uses explicit a:buNone, ordered list uses a:buAutoNum.",
        ],
    }
    (FIXTURE / "structure-report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"status": report["status"], "fixture": report["fixture"], "slideCount": len(rendered.slides), "report": str(FIXTURE / "structure-report.json")}, indent=2))


if __name__ == "__main__":
    main()
