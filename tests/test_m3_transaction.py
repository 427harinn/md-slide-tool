import json
import subprocess
import sys
import tempfile
import threading
import time
import unittest
import zipfile
from xml.etree import ElementTree as ET
from pathlib import Path
from unittest.mock import patch
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import m3_entry
from m3_transaction import (
    DocumentBusy, FileTransaction, ManualRecoveryRequired, _atomic_lock,
    bootstrap_key, recover_bootstrap_transactions, recover_transactions,
)
from m2_identity import atomic_write_json
from m3_core import write_custom_properties


class M3LockTests(unittest.TestCase):
    def test_bootstrap_recovery_is_scoped_to_its_qmd_owner(self):
        with tempfile.TemporaryDirectory() as temp:
            project = Path(temp)
            qmd_a = project / "a.qmd"
            qmd_b = project / "b.qmd"
            qmd_a.write_text("a", encoding="utf-8")
            qmd_b.write_text("b", encoding="utf-8")
            doc_a = "document-123e4567-e89b-42d3-a456-426614174000"
            doc_b = "document-123e4567-e89b-42d3-a456-426614174001"

            # A matching stale init directory may be safely cleaned once its
            # document lock is acquired.
            a_dir = project / ".md-slide-tool" / "transactions" / doc_a
            a_init = a_dir / f".stale.init.{bootstrap_key(qmd_a)}.dead"
            a_init.mkdir(parents=True)
            (a_init / "manifest.json").write_text("partial", encoding="utf-8")

            # An unrelated live first-render journal and init directory must
            # survive recovery for QMD A, even though both share a project.
            tx_b = FileTransaction(project, doc_b, "b-empty", paired_qmd_path=qmd_b)
            tx_b._save()
            b_init = tx_b.root.parent / f".b-live.init.{bootstrap_key(qmd_b)}.live"
            b_init.mkdir()
            (b_init / "manifest.json").write_text("partial", encoding="utf-8")

            recover_bootstrap_transactions(project, qmd_a)
            self.assertFalse(a_init.exists())
            self.assertTrue(tx_b.root.exists())
            self.assertTrue(b_init.exists())

    def test_bootstrap_recovery_does_not_touch_foreign_document_busy_state(self):
        with tempfile.TemporaryDirectory() as temp:
            project = Path(temp)
            qmd_a, qmd_b = project / "a.qmd", project / "b.qmd"
            qmd_a.write_text("a", encoding="utf-8")
            qmd_b.write_text("b", encoding="utf-8")
            doc_b = "document-123e4567-e89b-42d3-a456-426614174001"
            tx_b = FileTransaction(project, doc_b, "b-empty", paired_qmd_path=qmd_b)
            tx_b._save()
            lock = project / ".md-slide-tool" / "locks" / "document" / doc_b / "owner.lock"
            with _atomic_lock(lock):
                # Recovery for A must not try to acquire B's busy lock.
                recover_bootstrap_transactions(project, qmd_a)
            self.assertTrue(tx_b.root.exists())

    def test_legacy_asset_migration_reuses_paired_pptx_identity(self):
        with tempfile.TemporaryDirectory() as temp:
            project = Path(temp)
            qmd = project / "deck.qmd"
            qmd.write_text("---\ntitle: Deck\n---\n", encoding="utf-8")
            pptx = project / "deck.pptx"
            Presentation().save(pptx)
            document_id = "document-123e4567-e89b-42d3-a456-426614174000"
            write_custom_properties(pptx, {
                "md-slide-tool.managed": "1",
                "md-slide-tool.format-version": "1",
                "md-slide-tool.document-id": document_id,
            })
            assets = project / "deck.assets.json"
            assets.write_text(json.dumps({"schemaVersion": 1, "slideSize": None, "elements": []}), encoding="utf-8")
            self.assertEqual(m3_entry.bootstrap_document_id(qmd), document_id)
            self.assertEqual(json.loads(assets.read_text(encoding="utf-8"))["schemaVersion"], 1)

    def test_bootstrap_recovers_first_render_transaction_when_pair_has_no_id_yet(self):
        with tempfile.TemporaryDirectory() as temp:
            project = Path(temp)
            qmd = project / "deck.qmd"
            qmd.write_text("---\ntitle: Deck\n---\n", encoding="utf-8")
            pptx = project / "deck.pptx"
            Presentation().save(pptx)
            old_pptx = pptx.read_bytes()
            assets = project / "deck.assets.json"
            old_assets = json.dumps({"schemaVersion": 1, "slideSize": None, "elements": []}).encode()
            assets.write_bytes(old_assets)
            orphan_id = "document-123e4567-e89b-42d3-a456-426614174000"
            tx = FileTransaction(project, orphan_id, "interrupted-first-render")
            tx.stage(pptx, b"candidate-pptx")
            tx.stage(qmd, b"candidate-qmd")
            tx.stage(assets, b"candidate-assets")
            tx.manifest["phase"] = "committing"
            tx._save()
            observed = {}

            def fake_render(command, **_kwargs):
                observed["activeId"] = command[-1]
                observed["pptxRestored"] = pptx.read_bytes() == old_pptx
                observed["assetsRestored"] = assets.read_bytes() == old_assets
                atomic_write_json(assets, {"schemaVersion": 2, "documentId": command[-1],
                                           "slideSize": None, "elements": []})
                return 0

            with patch.object(m3_entry.subprocess, "call", fake_render):
                self.assertEqual(m3_entry.run_locked(qmd, ["render"]), 0)
            self.assertNotEqual(observed["activeId"], orphan_id)
            self.assertTrue(observed["pptxRestored"])
            self.assertTrue(observed["assetsRestored"])

    def test_bootstrap_to_document_handoff_releases_bootstrap_after_document_lock(self):
        with tempfile.TemporaryDirectory() as temp:
            project = Path(temp)
            qmd = project / "deck.qmd"
            qmd.write_text("---\ntitle: Deck\n---\n", encoding="utf-8")
            assets = project / "deck.assets.json"
            assets.write_text(json.dumps({"schemaVersion": 1, "slideSize": None, "elements": []}), encoding="utf-8")
            observed = {}

            def fake_render(command, **_kwargs):
                document_id = command[-1]
                observed["documentId"] = document_id
                locks = project / ".md-slide-tool" / "locks"
                self.assertEqual(json.loads(assets.read_text(encoding="utf-8"))["schemaVersion"], 1)
                # Both locks span candidate rendering and migration commit.
                with self.assertRaises(DocumentBusy):
                    with _atomic_lock(locks / "bootstrap" / m3_entry.bootstrap_key(qmd) / "owner.lock"):
                        pass
                try:
                    with _atomic_lock(locks / "document" / document_id / "owner.lock"):
                        observed["documentBusy"] = False
                except DocumentBusy:
                    observed["documentBusy"] = True
                atomic_write_json(assets, {"schemaVersion": 2, "documentId": document_id,
                                           "slideSize": None, "elements": []})
                return 0

            with patch.object(m3_entry.subprocess, "call", fake_render):
                self.assertEqual(m3_entry.run_locked(qmd, ["render"]), 0)
            self.assertEqual(json.loads(assets.read_text(encoding="utf-8"))["schemaVersion"], 2)
            self.assertTrue(observed["documentId"].startswith("document-"))
            self.assertTrue(observed["documentBusy"])

    def test_two_concurrent_first_renders_generate_one_id_and_second_gets_busy(self):
        with tempfile.TemporaryDirectory() as temp:
            project = Path(temp)
            qmd = project / "deck.qmd"
            qmd.write_text("---\ntitle: Deck\n---\n", encoding="utf-8")
            assets = project / "deck.assets.json"
            assets.write_text(json.dumps({"schemaVersion": 1, "slideSize": None, "elements": []}), encoding="utf-8")
            entered = threading.Event()
            release = threading.Event()
            outcomes = []

            def fake_render(_command, **_kwargs):
                document_id = _command[-1]
                atomic_write_json(assets, {"schemaVersion": 2, "documentId": document_id,
                                           "slideSize": None, "elements": []})
                entered.set()
                release.wait(5)
                return 0

            def first():
                try:
                    outcomes.append(("first", m3_entry.run_locked(qmd, ["render"])))
                except Exception as exc:
                    outcomes.append(("first-error", str(exc)))

            with patch.object(m3_entry.subprocess, "call", fake_render):
                worker = threading.Thread(target=first)
                worker.start()
                self.assertTrue(entered.wait(5))
                migrated_id = json.loads(assets.read_text(encoding="utf-8"))["documentId"]
                with self.assertRaises(DocumentBusy):
                    m3_entry.run_locked(qmd, ["render"])
                release.set()
                worker.join(5)
            self.assertFalse(worker.is_alive())
            self.assertEqual(outcomes, [("first", 0)])
            self.assertEqual(json.loads(assets.read_text(encoding="utf-8"))["documentId"], migrated_id)

    def test_bootstrap_reread_reuses_id_created_before_lock_acquisition(self):
        with tempfile.TemporaryDirectory() as temp:
            project = Path(temp)
            qmd = project / "deck.qmd"
            qmd.write_text("---\ntitle: Deck\n---\n", encoding="utf-8")
            doc_id = "document-123e4567-e89b-42d3-a456-426614174000"
            assets = project / "deck.assets.json"
            assets.write_text(json.dumps({"schemaVersion": 2, "documentId": doc_id, "slideSize": None, "elements": []}), encoding="utf-8")
            checks = {}

            def fake_render(_command, **_kwargs):
                locks = project / ".md-slide-tool" / "locks"
                checks["reused"] = json.loads(assets.read_text(encoding="utf-8"))["documentId"] == doc_id
                with self.assertRaises(DocumentBusy):
                    with _atomic_lock(locks / "bootstrap" / m3_entry.bootstrap_key(qmd) / "owner.lock"):
                        pass
                with self.assertRaises(DocumentBusy):
                    with _atomic_lock(locks / "document" / doc_id / "owner.lock"):
                        pass
                return 0

            with patch.object(m3_entry, "current_document_id", return_value=None), \
                 patch.object(m3_entry.subprocess, "call", fake_render):
                self.assertEqual(m3_entry.run_locked(qmd, ["render"]), 0)
            self.assertEqual(assets.read_text(encoding="utf-8"), json.dumps({"schemaVersion": 2, "documentId": doc_id, "slideSize": None, "elements": []}))
            self.assertTrue(checks["reused"])

    def test_migrated_import_is_document_locked_before_importer_runs(self):
        with tempfile.TemporaryDirectory() as temp:
            project = Path(temp)
            qmd = project / "deck.qmd"
            qmd.write_text("---\ntitle: Deck\n---\n", encoding="utf-8")
            doc_id = "document-123e4567-e89b-42d3-a456-426614174000"
            (project / "deck.assets.json").write_text(json.dumps({"schemaVersion": 2, "documentId": doc_id, "slideSize": None, "elements": []}), encoding="utf-8")
            called = []

            def fake_import(command, **_kwargs):
                called.append(command)
                lock = project / ".md-slide-tool" / "locks" / "document" / doc_id / "owner.lock"
                with self.assertRaises(DocumentBusy):
                    with _atomic_lock(lock):
                        pass
                return 0

            with patch.object(m3_entry.subprocess, "call", fake_import):
                self.assertEqual(m3_entry.run_locked(qmd, ["import", str(project / "deck.pptx")]), 0)
            self.assertEqual(len(called), 1)


class M3TransactionTests(unittest.TestCase):
    def test_custom_property_update_preserves_unmanaged_xml_types(self):
        with tempfile.TemporaryDirectory() as temp:
            pptx = Path(temp) / "typed.pptx"
            Presentation().save(pptx)
            write_custom_properties(pptx, {})
            custom = b'''<?xml version="1.0" encoding="UTF-8"?>
<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties"
 xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">
 <property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="2" name="Integer"><vt:i4>42</vt:i4></property>
 <property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="3" name="Boolean"><vt:bool>true</vt:bool></property>
 <property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="4" name="Date"><vt:filetime>2026-10-23T00:00:00Z</vt:filetime></property>
</Properties>'''
            tmp = pptx.with_suffix(".rewrite")
            with zipfile.ZipFile(pptx) as source, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dest:
                for name in source.namelist():
                    dest.writestr(name, custom if name == "docProps/custom.xml" else source.read(name))
            tmp.replace(pptx)
            write_custom_properties(pptx, {"md-slide-tool.managed": "1"})
            with zipfile.ZipFile(pptx) as package:
                root = ET.fromstring(package.read("docProps/custom.xml"))
            ns = {"cp": "http://schemas.openxmlformats.org/officeDocument/2006/custom-properties"}
            types = {node.attrib["name"]: next(iter(node)).tag.rsplit("}", 1)[-1]
                     for node in root.findall("cp:property", ns)}
            values = {node.attrib["name"]: next(iter(node)).text
                      for node in root.findall("cp:property", ns)}
            self.assertEqual(types["Integer"], "i4")
            self.assertEqual(types["Boolean"], "bool")
            self.assertEqual(types["Date"], "filetime")
            self.assertEqual((values["Integer"], values["Boolean"], values["Date"]),
                             ("42", "true", "2026-10-23T00:00:00Z"))

    def test_crash_before_initial_journal_directory_rename_is_discarded(self):
        with tempfile.TemporaryDirectory() as temp:
            project = Path(temp)
            target = project / "untouched.txt"
            target.write_bytes(b"old")
            tx_parent = project / ".md-slide-tool" / "transactions" / "document-x"
            init_dir = tx_parent / ".transaction.init.crash"
            init_dir.mkdir(parents=True)
            (init_dir / "manifest.json").write_text("partial", encoding="utf-8")
            recover_transactions(project, "document-x")
            self.assertEqual(target.read_bytes(), b"old")
            self.assertFalse(init_dir.exists())

    def test_crash_during_staging_leaves_discardable_prepared_journal(self):
        with tempfile.TemporaryDirectory() as temp:
            project = Path(temp)
            target = project / "a.txt"
            target.write_bytes(b"old")
            tx = FileTransaction(project, "document-x", "partial-stage")
            # Simulate death after materializing an orphan blob but before the
            # staged operation becomes part of the prepared journal.
            tx._save()
            (tx.root / "new").mkdir(parents=True)
            (tx.root / "new" / "orphan").write_bytes(b"uncommitted")
            recover_transactions(project, "document-x")
            self.assertEqual(target.read_bytes(), b"old")
            self.assertFalse(tx.root.exists())

    def test_prepared_transaction_is_discarded_without_mutation(self):
        with tempfile.TemporaryDirectory() as temp:
            project = Path(temp)
            target = project / "a.txt"
            target.write_bytes(b"old")
            tx = FileTransaction(project, "document-x", "prepared")
            tx.stage(target, b"new")
            tx._save()
            recover_transactions(project, "document-x")
            self.assertEqual(target.read_bytes(), b"old")
            self.assertFalse(tx.root.exists())

    def test_rolling_back_recovery_is_idempotent_after_partial_rollback(self):
        with tempfile.TemporaryDirectory() as temp:
            project = Path(temp)
            first, second = project / "first", project / "second"
            first.write_bytes(b"old-1")
            second.write_bytes(b"old-2")
            tx = FileTransaction(project, "document-x", "rolling")
            tx.stage(first, b"new-1")
            tx.stage(second, b"new-2")
            tx._save()
            tx.manifest["phase"] = "rolling_back"
            tx._save()
            (tx.root / "new" / "0").replace(first)
            # Recovery should restore the changed first file and leave the
            # untouched second file in its old state.
            recover_transactions(project, "document-x")
            self.assertEqual(first.read_bytes(), b"old-1")
            self.assertEqual(second.read_bytes(), b"old-2")

    def test_recovery_rejects_tampered_external_target_before_touching_it(self):
        with tempfile.TemporaryDirectory() as temp, tempfile.TemporaryDirectory() as outside:
            project = Path(temp)
            external = Path(outside) / "victim.txt"
            external.write_bytes(b"keep me")
            tx_root = project / ".md-slide-tool" / "transactions" / "document-x" / "tampered"
            tx_root.mkdir(parents=True)
            manifest = {
                "version": 1,
                "phase": "committing",
                "operations": [{
                    "target": str(external.resolve()), "operation": "create",
                    "existed": False, "oldHash": None,
                    "newHash": __import__("hashlib").sha256(b"malicious").hexdigest(),
                }],
            }
            (tx_root / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
            with self.assertRaisesRegex(ManualRecoveryRequired, "escapes project directory"):
                recover_transactions(project, "document-x")
            self.assertEqual(external.read_bytes(), b"keep me")


if __name__ == "__main__":
    unittest.main()
