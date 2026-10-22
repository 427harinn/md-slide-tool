import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
PREPARE = ROOT / "scripts" / "prepare-qmd-for-pptx.py"
POSTPROCESS = ROOT / "scripts" / "postprocess-pptx.py"


def run_prepare(qmd: Path, work: Path):
    return subprocess.run(
        [sys.executable, str(PREPARE), "--input", str(qmd), "--output", str(work / "render.qmd"), "--images-json", str(work / "images.json")],
        text=True, capture_output=True,
    )


def make_pptx(path: Path, title: str):
    run_python("""from pptx import Presentation
from pptx.util import Inches
import sys
prs = Presentation()
prs.slide_width = Inches(13.333333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])
heading = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(8), Inches(0.8))
heading.text = sys.argv[2]
prs.save(sys.argv[1])""", path, title)


def make_stable_blank_pptx(path: Path, slide_id: str):
    run_python("""from pptx import Presentation
from pptx.util import Inches
import sys
prs = Presentation()
prs.slide_width = Inches(13.333333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.name = sys.argv[2]
prs.save(sys.argv[1])""", path, slide_id)


def make_image(path: Path, size, color):
    run_python("""from PIL import Image
import sys
Image.new('RGB', (int(sys.argv[2]), int(sys.argv[3])), sys.argv[4]).save(sys.argv[1])""", path, size[0], size[1], color)


def run_python(code, *args):
    result = subprocess.run([sys.executable, "-c", code, *map(str, args)], text=True, capture_output=True)
    if result.returncode:
        raise AssertionError(result.stderr)
    return result.stdout


class StableIdsAssetsTests(unittest.TestCase):
    def test_missing_ids_are_added_and_existing_ids_stay_stable_without_content_loss(self):
        with tempfile.TemporaryDirectory() as directory:
            work = Path(directory)
            qmd = work / "deck.qmd"
            source = '---\ntitle: Deck\n---\n\n## Keep this title\n\nBody exactly as written.\n\n```md\n## Example heading\n![Example](example.png){element-id="not-a-real-id"}\n```\n\n![Alt caption](images/diagram.png){place="right"}\n'
            qmd.write_text(source, encoding="utf-8")
            first = run_prepare(qmd, work)
            self.assertEqual(first.returncode, 0, first.stderr)
            updated = qmd.read_text(encoding="utf-8")
            render_output = (work / "render.qmd").read_text(encoding="utf-8")
            self.assertIn("## Keep this title {slide-id=\"slide-", updated)
            self.assertIn('![Alt caption](images/diagram.png){place="right" element-id="element-', updated)
            self.assertIn("## Keep this title {slide-id=\"slide-", render_output)
            self.assertNotIn("images/diagram.png", render_output)
            self.assertIn("Body exactly as written.", updated)
            self.assertIn('![Example](example.png){element-id="not-a-real-id"}', updated)
            data = json.loads((work / "images.json").read_text(encoding="utf-8"))
            slide_id = data[0]["slide_id"]
            element_id = data[0]["images"][0]["element_id"]
            second = run_prepare(qmd, work)
            self.assertEqual(second.returncode, 0, second.stderr + second.stdout)
            second_data = json.loads((work / "images.json").read_text(encoding="utf-8"))
            self.assertEqual((second_data[0]["slide_id"], second_data[0]["images"][0]["element_id"]), (slide_id, element_id))

    def test_duplicate_slide_and_element_ids_fail_without_repair(self):
        for duplicate in (
            '## A {slide-id="slide-123e4567-e89b-42d3-a456-426614174000"}\n## B {slide-id="slide-123e4567-e89b-42d3-a456-426614174000"}\n',
            '## A {slide-id="slide-123e4567-e89b-42d3-a456-426614174000"}\n![one](a.png){element-id="element-123e4567-e89b-42d3-a456-426614174000"}\n![two](b.png){element-id="element-123e4567-e89b-42d3-a456-426614174000"}\n',
        ):
            with self.subTest(duplicate=duplicate), tempfile.TemporaryDirectory() as directory:
                work = Path(directory)
                qmd = work / "deck.qmd"
                qmd.write_text(duplicate, encoding="utf-8")
                result = run_prepare(qmd, work)
                self.assertNotEqual(result.returncode, 0)
                self.assertIn("duplicate", result.stderr)
                self.assertEqual(qmd.read_text(encoding="utf-8"), duplicate)

    def test_auto_layout_assets_restore_title_rename_and_path_change_by_ids(self):
        with tempfile.TemporaryDirectory() as directory:
            work = Path(directory)
            qmd = work / "deck.qmd"
            image_path = work / "diagram.png"
            make_image(image_path, (500, 300), "red")
            slide_id = "slide-123e4567-e89b-42d3-a456-426614174000"
            element_id = "element-123e4567-e89b-42d3-a456-426614174000"
            images = [{"slide_index": 0, "title": "Old title", "slide_id": slide_id, "images": [{"path": "diagram.png", "element_id": element_id, "caption": "caption", "place": "right"}]}]
            images_json = work / ".deck.images.json"
            assets = work / "deck.assets.json"
            images_json.write_text(json.dumps(images), encoding="utf-8")
            pptx = work / "deck.pptx"
            make_stable_blank_pptx(pptx, slide_id)
            first = subprocess.run([sys.executable, str(POSTPROCESS), "--pptx", str(pptx), "--images-json", str(images_json), "--assets", str(assets)], text=True, capture_output=True)
            self.assertEqual(first.returncode, 0, first.stderr)
            first_assets = json.loads(assets.read_text(encoding="utf-8"))
            self.assertEqual(first_assets["schemaVersion"], 1)
            self.assertEqual(first_assets["slideSize"]["unit"], "emu")
            entry = first_assets["elements"][0]
            self.assertEqual((entry["slideId"], entry["elementId"], entry["type"]), (slide_id, element_id, "image"))
            self.assertTrue(all(type(entry[key]) is int for key in ("x", "y", "width", "height")))
            saved_rect = tuple(entry[key] for key in ("x", "y", "width", "height"))

            # Simulate a fresh Quarto output with the renamed title and changed image path.
            replacement = work / "replacement.png"
            make_image(replacement, (500, 400), "blue")
            images[0]["title"] = "Renamed title"
            images[0]["images"][0]["path"] = "replacement.png"
            images_json.write_text(json.dumps(images), encoding="utf-8")
            make_stable_blank_pptx(pptx, slide_id)
            second = subprocess.run([sys.executable, str(POSTPROCESS), "--pptx", str(pptx), "--images-json", str(images_json), "--assets", str(assets)], text=True, capture_output=True)
            self.assertEqual(second.returncode, 0, second.stderr)
            reopened = json.loads(run_python("""from pptx import Presentation
import json, sys
prs = Presentation(sys.argv[1])
slide = prs.slides[0]
picture = next(shape for shape in slide.shapes if shape.shape_type == 13)
print(json.dumps({'slide': slide.name, 'picture': picture.name, 'rect': [int(picture.left), int(picture.top), int(picture.width), int(picture.height)]}))""", pptx))
            self.assertEqual(reopened["picture"], element_id)
            self.assertEqual(tuple(reopened["rect"]), saved_rect, second.stdout)
            self.assertEqual(reopened["slide"], slide_id)
            updated_assets = json.loads(assets.read_text(encoding="utf-8"))
            self.assertEqual(updated_assets["elements"][0]["path"], "replacement.png")
            self.assertNotIn("title", updated_assets["elements"][0])
            self.assertNotIn("caption", updated_assets["elements"][0])

    def test_invalid_layout_falls_back_and_stale_entries_are_pruned_after_success(self):
        with tempfile.TemporaryDirectory() as directory:
            work = Path(directory)
            make_image(work / "img.png", (500, 400), "green")
            slide_id = "slide-123e4567-e89b-42d3-a456-426614174000"
            element_id = "element-123e4567-e89b-42d3-a456-426614174000"
            images_json = work / "deck.images.json"
            assets = work / "deck.assets.json"
            images_json.write_text(json.dumps([{"title": "Slide", "slide_id": slide_id, "images": [{"path": "img.png", "element_id": element_id}]}]), encoding="utf-8")
            assets.write_text(json.dumps({"schemaVersion": 1, "slideSize": {"width": 1, "height": 1, "unit": "emu"}, "elements": [
                {"slideId": slide_id, "elementId": element_id, "type": "image", "path": "img.png", "x": "bad", "y": 0, "width": 1, "height": 1},
                {"slideId": slide_id, "elementId": "element-123e4567-e89b-42d3-a456-426614174001", "type": "image", "path": "stale.png", "x": 0, "y": 0, "width": 1, "height": 1},
            ]}), encoding="utf-8")
            pptx = work / "deck.pptx"
            make_pptx(pptx, "Slide")
            result = subprocess.run([sys.executable, str(POSTPROCESS), "--pptx", str(pptx), "--images-json", str(images_json), "--assets", str(assets)], text=True, capture_output=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertIn("[WARN] invalid saved layout", result.stdout)
            final = json.loads(assets.read_text(encoding="utf-8"))
            self.assertEqual(len(final["elements"]), 1)
            self.assertEqual(final["elements"][0]["elementId"], element_id)
            self.assertTrue(all(type(final["elements"][0][key]) is int for key in ("x", "y", "width", "height")))

    def test_existing_offslide_auto_layout_is_restored_exactly(self):
        with tempfile.TemporaryDirectory() as directory:
            work = Path(directory)
            make_image(work / "tiny.png", (40, 40), "purple")
            slide_id = "slide-123e4567-e89b-42d3-a456-426614174000"
            element_id = "element-123e4567-e89b-42d3-a456-426614174000"
            images_json = work / "deck.images.json"
            assets = work / "deck.assets.json"
            images_json.write_text(json.dumps([{"title": "Slide", "slide_id": slide_id, "images": [{"path": "tiny.png", "element_id": element_id}]}]), encoding="utf-8")
            pptx = work / "deck.pptx"
            make_stable_blank_pptx(pptx, slide_id)
            first = subprocess.run([sys.executable, str(POSTPROCESS), "--pptx", str(pptx), "--images-json", str(images_json), "--assets", str(assets)], text=True, capture_output=True)
            self.assertEqual(first.returncode, 0, first.stderr)
            entry = json.loads(assets.read_text(encoding="utf-8"))["elements"][0]
            first_rect = tuple(entry[key] for key in ("x", "y", "width", "height"))
            self.assertGreater(entry["x"], json.loads(assets.read_text(encoding="utf-8"))["slideSize"]["width"])

            make_stable_blank_pptx(pptx, slide_id)
            second = subprocess.run([sys.executable, str(POSTPROCESS), "--pptx", str(pptx), "--images-json", str(images_json), "--assets", str(assets)], text=True, capture_output=True)
            self.assertEqual(second.returncode, 0, second.stderr)
            result = json.loads(run_python("""from pptx import Presentation
import json, sys
prs = Presentation(sys.argv[1])
shape = next(item for item in prs.slides[0].shapes if item.shape_type == 13)
print(json.dumps([int(shape.left), int(shape.top), int(shape.width), int(shape.height)]))""", pptx))
            self.assertEqual(tuple(result), first_rect)

    def test_unbounded_out_of_slide_layout_warns_and_falls_back(self):
        for field, invalid_value in (("x", 10**20), ("y", 10**20), ("width", 10**20)):
            with self.subTest(field=field), tempfile.TemporaryDirectory() as directory:
                work = Path(directory)
                make_image(work / "tiny.png", (40, 40), "purple")
                slide_id = "slide-123e4567-e89b-42d3-a456-426614174000"
                element_id = "element-123e4567-e89b-42d3-a456-426614174000"
                images_json = work / "deck.images.json"
                images_json.write_text(json.dumps([{"title": "Slide", "slide_id": slide_id, "images": [{"path": "tiny.png", "element_id": element_id}]}]), encoding="utf-8")
                assets = work / "deck.assets.json"
                bad_entry = {"slideId": slide_id, "elementId": element_id, "type": "image", "path": "tiny.png", "x": 0, "y": 0, "width": 1, "height": 1}
                bad_entry[field] = invalid_value
                assets.write_text(json.dumps({"schemaVersion": 1, "slideSize": {"width": 1, "height": 1, "unit": "emu"}, "elements": [bad_entry]}), encoding="utf-8")
                pptx = work / "deck.pptx"
                make_stable_blank_pptx(pptx, slide_id)
                result = subprocess.run([sys.executable, str(POSTPROCESS), "--pptx", str(pptx), "--images-json", str(images_json), "--assets", str(assets)], text=True, capture_output=True)
                self.assertEqual(result.returncode, 0, result.stderr)
                self.assertIn("[WARN] invalid saved layout", result.stdout)
                final = json.loads(assets.read_text(encoding="utf-8"))["elements"][0]
                self.assertEqual(final["x"], json.loads(assets.read_text(encoding="utf-8"))["slideSize"]["width"] + int(0.08 * 914400))
                self.assertLessEqual(final["width"], int(2.5 * 914400))

    def test_selected_pptx_metadata_survives_save_reopen_roundtrip(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "roundtrip.pptx"
            make_pptx(path, "Slide")
            stable_slide_id = "slide-123e4567-e89b-42d3-a456-426614174000"
            stable_element_id = "element-123e4567-e89b-42d3-a456-426614174000"
            picture_path = Path(directory) / "pic.png"
            make_image(picture_path, (20, 20), "black")
            reopened = json.loads(run_python("""from pptx import Presentation
from pptx.util import Inches
import json, sys
path, image, slide_id, element_id = sys.argv[1:]
prs = Presentation(path)
slide = prs.slides[0]
slide.name = slide_id
picture = slide.shapes.add_picture(image, Inches(1), Inches(1), width=Inches(1), height=Inches(1))
picture.name = element_id
prs.save(path)
prs = Presentation(path)
slide = prs.slides[0]
picture = next(shape for shape in slide.shapes if shape.shape_type == 13)
print(json.dumps({'slide': slide.name, 'picture': picture.name}))""", path, picture_path, stable_slide_id, stable_element_id))
            self.assertEqual(reopened["slide"], stable_slide_id)
            self.assertEqual(reopened["picture"], stable_element_id)

    def test_failed_pptx_processing_does_not_modify_existing_assets(self):
        with tempfile.TemporaryDirectory() as directory:
            work = Path(directory)
            pptx = work / "broken.pptx"
            pptx.write_text("not a pptx", encoding="utf-8")
            images_json = work / "images.json"
            images_json.write_text("[]", encoding="utf-8")
            assets = work / "deck.assets.json"
            original = '{"schemaVersion":1,"slideSize":{"width":1,"height":1,"unit":"emu"},"elements":[]}\n'
            assets.write_text(original, encoding="utf-8")
            result = subprocess.run([sys.executable, str(POSTPROCESS), "--pptx", str(pptx), "--images-json", str(images_json), "--assets", str(assets)], text=True, capture_output=True)
            self.assertNotEqual(result.returncode, 0)
            self.assertEqual(assets.read_text(encoding="utf-8"), original)


if __name__ == "__main__":
    unittest.main()
