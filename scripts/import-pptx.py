#!/usr/bin/env python3
"""Strict M3 import of an md-slide-tool managed presentation."""

import argparse
import hashlib
import json
import re
import subprocess
import shutil
import sys
import string
import tempfile
import uuid
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from m2_identity import atomic_write_json, new_id, validate_id
from m3_core import (
    MDST_PROPS, qmd_body_sha256, read_custom_properties, split_front_matter,
    validate_assets, validate_crop, valid_document_id, write_custom_properties,
)
from m3_transaction import FileTransaction

from pptx import Presentation

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
PICTURE_SHAPE_TYPE = 13
TEXT_BOX_SHAPE_TYPE = 17


def shape_placeholder_type(shape):
    if not shape.is_placeholder:
        return None
    return int(shape.placeholder_format.type)


def ppr(paragraph):
    return paragraph._p.find(f"{{{A}}}pPr")


def has_marker(ppr_element, local_name):
    return ppr_element is not None and ppr_element.find(f"{{{A}}}{local_name}") is not None


def inherited_list_kind(shape, paragraph):
    """Resolve inherited list semantics via the placeholder/master style."""
    if not shape.is_placeholder:
        return None
    level = paragraph.level
    # Explicit body placeholder style, then master txStyles for this level.
    idx = int(shape.placeholder_format.idx)
    layout = shape.part.slide_layout
    for candidate in layout.placeholders:
        if int(candidate.placeholder_format.idx) == idx:
            body = candidate._element.find(f"{{{P}}}txBody")
            if body is not None:
                for p in body.findall(f"{{{A}}}p"):
                    pp = p.find(f"{{{A}}}pPr")
                    if pp is not None and int(pp.get("lvl", "0")) == level:
                        if pp.find(f"{{{A}}}buChar") is not None or pp.find(f"{{{A}}}buAutoNum") is not None:
                            return "numbered" if pp.find(f"{{{A}}}buAutoNum") is not None else "bullet"
                        if pp.find(f"{{{A}}}buNone") is not None:
                            return None
    master = layout.slide_master._element
    txstyles = master.find(f"{{{P}}}txStyles")
    body_style = txstyles.find(f"{{{P}}}bodyStyle") if txstyles is not None else None
    if body_style is not None:
        lvl = body_style.find(f"{{{A}}}lvl{level + 1}pPr")
        if lvl is not None:
            if lvl.find(f"{{{A}}}buAutoNum") is not None:
                return "numbered"
            if lvl.find(f"{{{A}}}buChar") is not None:
                return "bullet"
    return None


def is_title_layout(slide):
    """Recognize supported title layouts from their placeholder structure."""
    layout_types = {shape_placeholder_type(shape) for shape in slide.slide_layout.shapes}
    # Both repository fixtures use one of these structural title/subtitle
    # pairs. Names, geometry, and colors are deliberately irrelevant.
    return (1 in layout_types or 3 in layout_types) and 4 in layout_types


def text_and_kind(shape):
    items = []
    for paragraph in shape.text_frame.paragraphs:
        if paragraph.level != 0:
            raise ValueError("nested list levels are unsupported")
        allowed_children = {f"{{{A}}}pPr", f"{{{A}}}endParaRPr", f"{{{A}}}r"}
        if any(child.tag not in allowed_children for child in paragraph._p):
            raise ValueError("unsupported PowerPoint paragraph structure")
        for run in paragraph.runs:
            allowed_run_children = {f"{{{A}}}rPr", f"{{{A}}}t"}
            if any(child.tag not in allowed_run_children for child in run._r):
                raise ValueError("unsupported PowerPoint text run structure")
        text = paragraph.text.strip()
        if not text:
            continue
        pp = ppr(paragraph)
        if has_marker(pp, "buNone"):
            kind = "paragraph"
        elif has_marker(pp, "buAutoNum"):
            kind = "numbered"
        elif has_marker(pp, "buChar"):
            kind = "bullet"
        else:
            # Repository body placeholders inherit bullets at lvl 0. An
            # unmarked paragraph with no inherited list style is ordinary text.
            kind = inherited_list_kind(shape, paragraph) or "paragraph"
        items.append((kind, text))
    return items


def escape_markdown(text):
    # Imported strings are literal PowerPoint text. Escape every ASCII
    # punctuation character so punctuation cannot become Markdown syntax,
    # raw HTML, or a block/list delimiter when reconstructed into QMD.
    text = text.replace("\n", " ").replace("\r", " ")
    return "".join(("\\" + char) if char in string.punctuation else char for char in text)


def crop_for_picture(shape):
    src = shape._element.xpath("./p:blipFill/a:srcRect")
    if not src:
        return None
    node = src[0]
    crop = {"unit": "ooxml-percent"}
    for key, xml_key in (("left", "l"), ("top", "t"), ("right", "r"), ("bottom", "b")):
        try:
            crop[key] = int(node.get(xml_key, "0"))
        except ValueError as exc:
            raise ValueError("invalid non-integer PowerPoint crop") from exc
    return validate_crop(crop)


def layout_crop(picture):
    return crop_for_picture(picture)


def slide_is_title_role(slide, document_id):
    expected = f"mdst-title-{document_id.removeprefix('document-')}"
    return slide.name == expected


def reconstruct(pptx_path: Path, qmd_path: Path, assets_path: Path, tx: FileTransaction):
    props = read_custom_properties(pptx_path)
    if props.get("md-slide-tool.managed") != "1" or props.get("md-slide-tool.format-version") != "1":
        raise ValueError("unmanaged/unsupported PPTX; render with an M3-capable version first")
    document_id = props.get("md-slide-tool.document-id")
    if not valid_document_id(document_id):
        raise ValueError("invalid PPTX document-id")
    raw_qmd = qmd_path.read_bytes()
    expected_hash = props.get("md-slide-tool.qmd-body-sha256")
    if not re.fullmatch(r"[0-9a-f]{64}", expected_hash or ""):
        raise ValueError("invalid QMD baseline metadata")
    if qmd_body_sha256(raw_qmd) != expected_hash:
        raise ValueError("QMD_CHANGED_SINCE_RENDER")
    front, _ = split_front_matter(raw_qmd)
    try:
        prior = json.loads(assets_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ValueError("same-stem assets file is missing or invalid") from exc
    assets = validate_assets(prior, document_id)
    if assets["schemaVersion"] != 2:
        raise ValueError("managed import requires assets schema v2; render first")
    prior_by_id = {entry["elementId"]: entry for entry in assets["elements"]}

    prs = Presentation(str(pptx_path))
    title_count = sum(slide_is_title_role(slide, document_id) for slide in prs.slides)
    if title_count > 1:
        raise ValueError("duplicate title-role marker")
    slide_ids = set()
    element_ids = set()
    caption_by_element = {}
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.name.startswith("mdst-caption-"):
                if shape.is_placeholder or int(shape.shape_type) != TEXT_BOX_SHAPE_TYPE:
                    raise ValueError("caption marker must be a non-placeholder text box")
                eid = shape.name[len("mdst-caption-"):]
                if not validate_id(eid, "element"):
                    raise ValueError("malformed caption marker")
                if eid in caption_by_element:
                    raise ValueError("duplicate caption marker")
                caption_by_element[eid] = shape.text.strip()

    rebuilt = []
    layout_entries = []
    extracted_writes = {}
    referenced_caption_ids = set()
    normal_slide_order = []
    prior_path_users = {}
    for entry in assets["elements"]:
        prior_path_users.setdefault(entry["path"], set()).add(entry["elementId"])
    for slide in prs.slides:
        if slide_is_title_role(slide, document_id):
            # The reserved title role is excluded from QMD reconstruction,
            # but its exact generated shape structure is still validated.
            # T-301 render fixtures emit exactly one center-title placeholder;
            # any subtitle/duplicate/other local shape may contain user data.
            if (len(slide.shapes) != 1 or
                    shape_placeholder_type(slide.shapes[0]) != 3):
                raise ValueError("unsupported slide-local structure on title slide")
            continue
        sid = slide.name
        title_shapes = [s for s in slide.shapes if shape_placeholder_type(s) in (1, 3)]
        title_layout = is_title_layout(slide)
        if not sid or not sid.startswith("slide-"):
            if title_layout:
                raise ValueError("ID-less title-layout slide is ambiguous")
            sid = new_id("slide")
            slide.name = sid
        if not validate_id(sid, "slide") or sid in slide_ids:
            raise ValueError("malformed or duplicate slide-id")
        slide_ids.add(sid)
        if title_layout:
            raise ValueError("unmarked title-layout slide is ambiguous")
        if len(title_shapes) != 1:
            raise ValueError("normal slide must have exactly one supported title placeholder")
        title = title_shapes[0].text.strip()
        if not title:
            raise ValueError("slide title is empty")
        body_shapes = [s for s in slide.shapes if shape_placeholder_type(s) in (2, 7)]
        if len(body_shapes) > 1:
            raise ValueError("multiple body placeholders are unsupported")
        contents = []
        if body_shapes:
            contents.extend(text_and_kind(body_shapes[0]))
        image_lines = []
        for shape in slide.shapes:
            phtype = shape_placeholder_type(shape)
            if phtype in (1, 2, 7):
                continue
            if phtype == 3:
                continue
            if shape.name.startswith("mdst-caption-"):
                continue
            if int(shape.shape_type) == PICTURE_SHAPE_TYPE:
                if int(shape.rotation) != 0:
                    raise ValueError("picture rotation is unsupported")
                eid = shape.name
                if eid.startswith("element-") and not validate_id(eid, "element"):
                    raise ValueError("malformed element-id")
                if not eid.startswith("element-"):
                    eid = new_id("element")
                    shape.name = eid
                if eid in element_ids:
                    raise ValueError("duplicate element-id")
                element_ids.add(eid)
                caption = caption_by_element.get(eid, "")
                if eid in caption_by_element:
                    referenced_caption_ids.add(eid)
                blob = shape.image.blob
                digest = hashlib.sha256(blob).hexdigest()
                previous = prior_by_id.get(eid)
                source = previous.get("path") if previous else None
                source_path = (qmd_path.parent / source).resolve(strict=False) if source else None
                if source_path and source_path.is_file() and hashlib.sha256(source_path.read_bytes()).hexdigest() == digest:
                    image_path = source
                else:
                    ext = shape.image.ext.lower() or "bin"
                    image_path = f"img/imported/{eid}.{ext}"
                    shared_by_other = prior_path_users.get(image_path, set()) - {eid}
                    if shared_by_other:
                        image_path = f"img/imported/{eid}-{digest[:12]}.{ext}"
                    extracted_writes[qmd_path.parent / image_path] = blob
                crop = layout_crop(shape)
                entry = {"slideId": sid, "elementId": eid, "type": "image", "path": image_path,
                         "x": int(shape.left), "y": int(shape.top), "width": int(shape.width), "height": int(shape.height)}
                if crop is not None:
                    entry["crop"] = crop
                layout_entries.append(entry)
                image_lines.append(f'![{escape_markdown(caption)}]({image_path}){{element-id="{eid}"}}')
                continue
            raise ValueError(f"unsupported slide-local shape: {shape.shape_type} / {shape.name}")
        lines = [f'## {escape_markdown(title)} {{slide-id="{sid}"}}']
        current_list = None
        for kind, text in contents:
            if kind == "paragraph":
                current_list = None
                lines.extend(["", escape_markdown(text)])
            else:
                if current_list != kind:
                    lines.append("")
                    current_list = kind
                lines.append(("- " if kind == "bullet" else "1. ") + escape_markdown(text))
        if image_lines:
            lines.extend(["", *image_lines])
        rebuilt.append("\n".join(lines))
        normal_slide_order.append(sid)
    orphaned = set(caption_by_element) - referenced_caption_ids
    if orphaned:
        raise ValueError("orphan caption marker")
    body = ("\n\n" if rebuilt else "").join(rebuilt)
    new_qmd = front + body.encode("utf-8")

    # Drop deleted images from imported storage only; never delete source files.
    present_ids = set(element_ids)
    retained_paths = {entry["path"] for entry in layout_entries}
    staged_deletes = set()
    for prior_entry in assets["elements"]:
        if prior_entry["elementId"] not in present_ids:
            old_path = prior_entry.get("path", "")
            if old_path.startswith("img/imported/") and old_path not in retained_paths and old_path not in staged_deletes:
                target = qmd_path.parent / old_path
                tx.stage(target, None)
                staged_deletes.add(old_path)

    new_assets = {"schemaVersion": 2, "documentId": document_id,
                  "slideSize": {"width": int(prs.slide_width), "height": int(prs.slide_height), "unit": "emu"},
                  "elements": layout_entries}
    # IDs and baseline are written into the staged PPTX before commit.
    with tempfile.TemporaryDirectory(dir=str(qmd_path.parent)) as tempdir:
        staged_pptx = Path(tempdir) / pptx_path.name
        shutil.copy2(pptx_path, staged_pptx)
        prs.save(str(staged_pptx))
        write_custom_properties(staged_pptx, {"md-slide-tool.qmd-body-sha256": qmd_body_sha256(new_qmd)})
        tx.stage(pptx_path, staged_pptx.read_bytes())
    tx.stage(qmd_path, new_qmd)
    tx.stage(assets_path, (json.dumps(new_assets, ensure_ascii=False, indent=2) + "\n").encode("utf-8"))
    for target, blob in extracted_writes.items():
        tx.stage(target, blob)
    return len(rebuilt), len(layout_entries)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--pptx", required=True)
    parser.add_argument("--locked", action="store_true", help=argparse.SUPPRESS)
    args = parser.parse_args()
    pptx_path = Path(args.pptx).resolve(strict=True)
    if not args.locked:
        entry = Path(__file__).resolve().parent / "m3_entry.py"
        return subprocess.call([sys.executable, str(entry), "import", str(pptx_path)])
    qmd_path = pptx_path.with_suffix(".qmd")
    assets_path = pptx_path.with_suffix(".assets.json")
    if not qmd_path.is_file():
        raise ValueError("same-stem QMD is required")
    props = read_custom_properties(pptx_path)
    doc_id = props.get("md-slide-tool.document-id")
    if not valid_document_id(doc_id):
        raise ValueError("unmanaged PPTX; render with an M3-capable version first")
    tx = FileTransaction(qmd_path.parent, doc_id, paired_qmd_path=qmd_path)
    slides, images = reconstruct(pptx_path, qmd_path, assets_path, tx)
    tx.commit()
    print(f"Imported {slides} slides and {images} images from {pptx_path}")


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise SystemExit(1)
